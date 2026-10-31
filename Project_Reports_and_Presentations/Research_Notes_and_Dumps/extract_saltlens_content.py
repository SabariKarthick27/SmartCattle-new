import docx
import json

doc = docx.Document('SaltLens_DSS_Final_Project_Report_Final_With_Cover_ORIGINAL_BACKUP.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

# Save table contents to a json file for precise retrieval
tables_data = []
for t_idx, t in enumerate(doc.tables):
    rows_data = []
    for r in t.rows:
        rows_data.append([c.text.strip() for c in r.cells])
    tables_data.append({
        'table_index': t_idx,
        'rows_count': len(t.rows),
        'cols_count': len(t.columns),
        'data': rows_data
    })

with open('saltlens_tables_extracted.json', 'w', encoding='utf-8') as f:
    json.dump(tables_data, f, indent=2, ensure_ascii=False)

print(f"Extracted {len(tables_data)} tables to saltlens_tables_extracted.json")
