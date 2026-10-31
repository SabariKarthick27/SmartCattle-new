import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def build_strict_saltlens_report():
    print("Beginning strict report build for SaltLens DSS...")
    doc = docx.Document()

    # Academic dimensions: A4 size
    # Margins: Left 1.25 in, Right 1.0 in, Top 1.0 in, Bottom 1.0 in
    TAB_STOP_RIGHT = 8640  # 6.0 in right margin tab

    # Setup Section 0: Cover Page & Title Page
    sec0 = doc.sections[0]
    sec0.top_margin = Inches(1.0)
    sec0.bottom_margin = Inches(1.0)
    sec0.left_margin = Inches(1.25)
    sec0.right_margin = Inches(1.0)
    sec0.page_width = Inches(8.27)
    sec0.page_height = Inches(11.69)

    COLOR_BLACK = RGBColor(0x00, 0x00, 0x00)
    COLOR_DARK = RGBColor(0x22, 0x22, 0x22)

    HEX_HDR_BG = "F2F4F7"
    HEX_ALT_BG = "FAFAFA"
    HEX_BORDER = "999999"

    # Base style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = COLOR_BLACK

    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=70, bottom=70, left=100, right=100):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def prevent_row_split(row):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    def set_repeat_header(row):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    def set_table_borders(table, color=HEX_BORDER, sz="4", val="single"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def add_page_number_to_footer(section):
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.text = ""
        r = p.add_run()
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        r.font.color.rgb = COLOR_BLACK
        fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
        instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
        fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
        fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
        r._r.append(fldChar1)
        r._r.append(instrText)
        r._r.append(fldChar2)
        r._r.append(fldChar3)

    print("Helper functions initialized successfully!")

if __name__ == '__main__':
    build_strict_saltlens_report()
