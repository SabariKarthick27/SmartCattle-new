import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import json

def build_full_strict_report():
    print("=" * 70)
    print("Building Full Strict SaltLens Report according to Sona / Anna Univ Template")
    print("=" * 70)

    doc = docx.Document()

    # Load JSON tables
    with open('saltlens_tables_extracted.json', 'r', encoding='utf-8') as f:
        all_tables_json = json.load(f)

    TAB_STOP_RIGHT = 8640  # 6.0 in right tab stop for dot leaders
    TAB_STOP_ABBR = 2160   # 1.5 in tab stop

    # Configure Section 0 (Cover & Title Page - No headers/footers, no page numbers)
    sec0 = doc.sections[0]
    sec0.page_width = Inches(8.27)
    sec0.page_height = Inches(11.69)
    sec0.top_margin = Inches(1.0)
    sec0.bottom_margin = Inches(1.0)
    sec0.left_margin = Inches(1.25)
    sec0.right_margin = Inches(1.0)

    COLOR_BLACK = RGBColor(0x00, 0x00, 0x00)
    HEX_HDR_BG = "EFEFEF"
    HEX_ALT_BG = "FAFAFA"
    HEX_BORDER = "888888"

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = COLOR_BLACK

    # XML Helper Functions
    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=70, bottom=70, left=100, right=100):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def prevent_row_split(row):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    def set_repeat_header(row):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    def set_table_borders(table, color=HEX_BORDER, sz="4", val="single"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def add_page_number_to_footer(section):
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.text = ""
        r = p.add_run()
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        r.font.color.rgb = COLOR_BLACK
        fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
        instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
        fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
        fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
        r._r.append(fldChar1)
        r._r.append(instrText)
        r._r.append(fldChar2)
        r._r.append(fldChar3)

    # Headings and Typography
    def add_chapter_heading(title_text, chapter_number=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        if chapter_number is not None:
            r1 = p.add_run(f"CHAPTER {chapter_number}\n")
            r1.font.name = 'Times New Roman'
            r1.font.size = Pt(15)
            r1.bold = True
            r1.font.color.rgb = COLOR_BLACK
            
        r2 = p.add_run(title_text.upper())
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(14)
        r2.bold = True
        r2.font.color.rgb = COLOR_BLACK
        return p

    def add_sec_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13)
        r.bold = True
        r.font.color.rgb = COLOR_BLACK
        return p

    def add_sec_heading_3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.bold = True
        r.font.color.rgb = COLOR_BLACK
        return p

    def add_sec_heading_4(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11.5)
        r.bold = True
        r.font.color.rgb = COLOR_BLACK
        return p

    def add_body_p(text, bold_prefix="", space_after=6, space_before=0, line_spacing=1.5, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = 'Times New Roman'
            r_pre.font.size = Pt(12)
            r_pre.bold = True
            r_pre.font.color.rgb = COLOR_BLACK
            
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.italic = italic
        r.font.color.rgb = COLOR_BLACK
        return p

    def add_bullet_p(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.35
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = 'Times New Roman'
            r_pre.font.size = Pt(12)
            r_pre.bold = True
            r_pre.font.color.rgb = COLOR_BLACK
            
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.font.color.rgb = COLOR_BLACK
        return p

    def add_figure_item(image_path, caption_text, width_inches=5.2):
        if os.path.exists(image_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(12)
            p_img.paragraph_format.space_after = Pt(4)
            p_img.paragraph_format.keep_with_next = True
            run_img = p_img.add_run()
            run_img.add_picture(image_path, width=Inches(width_inches))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(14)
            r_cap = p_cap.add_run(caption_text)
            r_cap.font.name = 'Times New Roman'
            r_cap.font.size = Pt(10.5)
            r_cap.bold = True
            r_cap.font.color.rgb = COLOR_BLACK

    def add_normal_table(headers, data, col_widths, caption_text=""):
        if caption_text:
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(14)
            p_cap.paragraph_format.space_after = Pt(4)
            p_cap.paragraph_format.keep_with_next = True
            r_cap = p_cap.add_run(caption_text)
            r_cap.font.name = 'Times New Roman'
            r_cap.font.size = Pt(10.5)
            r_cap.bold = True
            r_cap.font.color.rgb = COLOR_BLACK

        table = doc.add_table(rows=len(data) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, color=HEX_BORDER, sz="4", val="single")

        hdr_row = table.rows[0]
        set_repeat_header(hdr_row)
        prevent_row_split(hdr_row)

        for i, title in enumerate(headers):
            cell = hdr_row.cells[i]
            cell.width = col_widths[i]
            set_cell_background(cell, HEX_HDR_BG)
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(title)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10)
            r.bold = True
            r.font.color.rgb = COLOR_BLACK

        for r_idx, row_data in enumerate(data):
            row = table.rows[r_idx + 1]
            prevent_row_split(row)
            bg_color = HEX_ALT_BG if r_idx % 2 == 1 else "FFFFFF"

            for c_idx, val in enumerate(row_data):
                cell = row.cells[c_idx]
                cell.width = col_widths[c_idx]
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.2
                
                if c_idx == 0 and len(headers) > 2 and len(val) < 20:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif "%" in val or val.replace(".", "").isdigit() or "mm" in val or "cm" in val or "g/L" in val:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

                r = p.add_run(val)
                r.font.name = 'Times New Roman'
                r.font.size = Pt(9.5)
                r.font.color.rgb = COLOR_BLACK

        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_before = Pt(1)
        p_after.paragraph_format.space_after = Pt(8)

    def add_dot_leader_line(title_text, page_str, is_bold=False, indent_spaces=0):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.3
        pPr = p._p.get_or_add_pPr()
        tabs = parse_xml(f'<w:tabs {nsdecls("w")}><w:tab w:val="right" w:leader="dot" w:pos="{TAB_STOP_RIGHT}"/></w:tabs>')
        pPr.append(tabs)

        prefix = "    " * indent_spaces
        r1 = p.add_run(prefix + title_text)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(11)
        if is_bold:
            r1.bold = True
        r1.font.color.rgb = COLOR_BLACK

        r2 = p.add_run(f"\t{page_str}")
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(11)
        if is_bold:
            r2.bold = True
        r2.font.color.rgb = COLOR_BLACK
        return p

    def add_dot_leader_header(left_title="TITLE", right_title="PAGE NO."):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        pPr = p._p.get_or_add_pPr()
        tabs = parse_xml(f'<w:tabs {nsdecls("w")}><w:tab w:val="right" w:pos="{TAB_STOP_RIGHT}"/></w:tabs>')
        pPr.append(tabs)

        r1 = p.add_run(left_title)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(11.5)
        r1.bold = True
        r1.font.color.rgb = COLOR_BLACK

        r2 = p.add_run(f"\t{right_title}")
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(11.5)
        r2.bold = True
        r2.font.color.rgb = COLOR_BLACK

    # ==========================================
    # 1. COVER PAGE (Appendix 1 Specimen Format)
    # ==========================================
    print("Generating Cover Page...")
    # Logos Table at Top
    tbl_logos = doc.add_table(rows=1, cols=3)
    tbl_logos.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_logos, color="FFFFFF", sz="0", val="none")
    prevent_row_split(tbl_logos.rows[0])
    c_l, c_m, c_r = tbl_logos.rows[0].cells
    c_l.width = Inches(1.15)
    c_m.width = Inches(3.72)
    c_r.width = Inches(1.15)

    p_l = c_l.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists("extracted_header_logos/col_0_image1.jpeg"):
        p_l.add_run().add_picture("extracted_header_logos/col_0_image1.jpeg", width=Inches(1.05))

    p_m = c_m.paragraphs[0]
    p_m.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_m.paragraph_format.line_spacing = 1.25
    r_t = p_m.add_run("SALTLENS DSS: AN EVIDENCE-GATED DIGITAL TWIN FOR CLIMATE-RESILIENT SALT-PAN MANAGEMENT")
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(13)
    r_t.bold = True
    r_t.font.color.rgb = COLOR_BLACK

    p_r = c_r.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists("extracted_header_logos/col_2_image2.jpeg"):
        p_r.add_run().add_picture("extracted_header_logos/col_2_image2.jpeg", width=Inches(1.15))

    # Spacing and Details
    p_rep = doc.add_paragraph()
    p_rep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_rep.paragraph_format.space_before = Pt(38)
    p_rep.paragraph_format.space_after = Pt(4)
    r = p_rep.add_run("A Project Report")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True
    r.italic = True
    r.font.color.rgb = COLOR_BLACK

    p_course = doc.add_paragraph()
    p_course.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_course.paragraph_format.space_after = Pt(28)
    r = p_course.add_run("U23IT705 – PROJECT WORK")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(10)
    r = p_sub.add_run("Submitted by")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12.5)
    r.bold = True
    r.italic = True
    r.font.color.rgb = COLOR_BLACK

    # Students Table
    tbl_cand = doc.add_table(rows=2, cols=2)
    tbl_cand.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_cand, color="FFFFFF", sz="0", val="none")
    candidates = [
        ("SACHIN S B", "(61782323110095)"),
        ("THRISHNASANJAY S", "(61782323110117)")
    ]
    for row_idx, (name, regno) in enumerate(candidates):
        r_c = tbl_cand.rows[row_idx]
        prevent_row_split(r_c)
        c0, c1 = r_c.cells
        c0.width = Inches(2.8)
        c1.width = Inches(2.2)

        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p0.paragraph_format.space_after = Pt(4)
        run0 = p0.add_run(name)
        run0.font.name = 'Times New Roman'
        run0.font.size = Pt(12.5)
        run0.bold = True
        run0.font.color.rgb = COLOR_BLACK

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p1.paragraph_format.space_after = Pt(4)
        run1 = p1.add_run(f"   {regno}")
        run1.font.name = 'Times New Roman'
        run1.font.size = Pt(12.5)
        run1.bold = True
        run1.font.color.rgb = COLOR_BLACK

    p_part = doc.add_paragraph()
    p_part.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_part.paragraph_format.space_before = Pt(32)
    p_part.paragraph_format.space_after = Pt(4)
    r = p_part.add_run("in partial fulfillment for the award of the degree\nof")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.color.rgb = COLOR_BLACK

    p_deg = doc.add_paragraph()
    p_deg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_deg.paragraph_format.space_after = Pt(3)
    r = p_deg.add_run("BACHELOR OF ENGINEERING\nin\nINFORMATION TECHNOLOGY")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13.5)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(28)
    p_inst.paragraph_format.space_after = Pt(2)
    r = p_inst.add_run("SONA COLLEGE OF TECHNOLOGY (AUTONOMOUS), SALEM – 636005")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    p_univ = doc.add_paragraph()
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_univ.paragraph_format.space_after = Pt(2)
    r = p_univ.add_run("ANNA UNIVERSITY: CHENNAI 600 025")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12.5)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_date.paragraph_format.space_after = Pt(0)
    r = p_date.add_run("OCTOBER 2026")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    doc.add_page_break()

    # ==========================================
    # 2. TITLE PAGE (Identical text per Appendix 1)
    # ==========================================
    print("Generating Title Page...")
    tbl_logos2 = doc.add_table(rows=1, cols=3)
    tbl_logos2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_logos2, color="FFFFFF", sz="0", val="none")
    prevent_row_split(tbl_logos2.rows[0])
    c_l2, c_m2, c_r2 = tbl_logos2.rows[0].cells
    c_l2.width = Inches(1.15)
    c_m2.width = Inches(3.72)
    c_r2.width = Inches(1.15)

    p_l2 = c_l2.paragraphs[0]
    p_l2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists("extracted_header_logos/col_0_image1.jpeg"):
        p_l2.add_run().add_picture("extracted_header_logos/col_0_image1.jpeg", width=Inches(1.05))

    p_m2 = c_m2.paragraphs[0]
    p_m2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_m2.paragraph_format.line_spacing = 1.25
    r_t2 = p_m2.add_run("SALTLENS DSS: AN EVIDENCE-GATED DIGITAL TWIN FOR CLIMATE-RESILIENT SALT-PAN MANAGEMENT")
    r_t2.font.name = 'Times New Roman'
    r_t2.font.size = Pt(13)
    r_t2.bold = True
    r_t2.font.color.rgb = COLOR_BLACK

    p_r2 = c_r2.paragraphs[0]
    p_r2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists("extracted_header_logos/col_2_image2.jpeg"):
        p_r2.add_run().add_picture("extracted_header_logos/col_2_image2.jpeg", width=Inches(1.15))

    p_rep2 = doc.add_paragraph()
    p_rep2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_rep2.paragraph_format.space_before = Pt(38)
    p_rep2.paragraph_format.space_after = Pt(4)
    r = p_rep2.add_run("A Project Report")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True
    r.italic = True
    r.font.color.rgb = COLOR_BLACK

    p_course2 = doc.add_paragraph()
    p_course2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_course2.paragraph_format.space_after = Pt(28)
    r = p_course2.add_run("U23IT705 – PROJECT WORK")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    p_sub2 = doc.add_paragraph()
    p_sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub2.paragraph_format.space_after = Pt(10)
    r = p_sub2.add_run("Submitted by")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12.5)
    r.bold = True
    r.italic = True
    r.font.color.rgb = COLOR_BLACK

    tbl_cand2 = doc.add_table(rows=2, cols=2)
    tbl_cand2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_cand2, color="FFFFFF", sz="0", val="none")
    for row_idx, (name, regno) in enumerate(candidates):
        r_c = tbl_cand2.rows[row_idx]
        prevent_row_split(r_c)
        c0, c1 = r_c.cells
        c0.width = Inches(2.8)
        c1.width = Inches(2.2)

        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p0.paragraph_format.space_after = Pt(4)
        run0 = p0.add_run(name)
        run0.font.name = 'Times New Roman'
        run0.font.size = Pt(12.5)
        run0.bold = True
        run0.font.color.rgb = COLOR_BLACK

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p1.paragraph_format.space_after = Pt(4)
        run1 = p1.add_run(f"   {regno}")
        run1.font.name = 'Times New Roman'
        run1.font.size = Pt(12.5)
        run1.bold = True
        run1.font.color.rgb = COLOR_BLACK

    p_part2 = doc.add_paragraph()
    p_part2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_part2.paragraph_format.space_before = Pt(32)
    p_part2.paragraph_format.space_after = Pt(4)
    r = p_part2.add_run("in partial fulfillment for the award of the degree\nof")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.color.rgb = COLOR_BLACK

    p_deg2 = doc.add_paragraph()
    p_deg2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_deg2.paragraph_format.space_after = Pt(3)
    r = p_deg2.add_run("BACHELOR OF ENGINEERING\nin\nINFORMATION TECHNOLOGY")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13.5)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    p_inst2 = doc.add_paragraph()
    p_inst2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst2.paragraph_format.space_before = Pt(28)
    p_inst2.paragraph_format.space_after = Pt(2)
    r = p_inst2.add_run("SONA COLLEGE OF TECHNOLOGY (AUTONOMOUS), SALEM – 636005")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    p_univ2 = doc.add_paragraph()
    p_univ2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_univ2.paragraph_format.space_after = Pt(2)
    r = p_univ2.add_run("ANNA UNIVERSITY: CHENNAI 600 025")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12.5)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    p_date2 = doc.add_paragraph()
    p_date2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_date2.paragraph_format.space_after = Pt(0)
    r = p_date2.add_run("OCTOBER 2026")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    doc.add_page_break()

    # ==========================================
    # 3. BONAFIDE CERTIFICATE (Appendix 2 Format)
    # ==========================================
    print("Generating Bonafide Certificate...")
    p_b_hdr = doc.add_paragraph()
    p_b_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_b_hdr.paragraph_format.space_before = Pt(10)
    p_b_hdr.paragraph_format.space_after = Pt(2)
    r = p_b_hdr.add_run("SONA COLLEGE OF TECHNOLOGY (AUTONOMOUS)\nANNA UNIVERSITY: CHENNAI – 600 025")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    p_b_tit = doc.add_paragraph()
    p_b_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_b_tit.paragraph_format.space_before = Pt(16)
    p_b_tit.paragraph_format.space_after = Pt(22)
    r = p_b_tit.add_run("BONAFIDE CERTIFICATE")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    # The Bonafide Certificate shall be in double line spacing using Font Style Times New Roman and Font Size 14
    p_b_body = doc.add_paragraph()
    p_b_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_b_body.paragraph_format.line_spacing = 2.0
    p_b_body.paragraph_format.space_after = Pt(36)
    r_body = p_b_body.add_run(
        "Certified that this project titled “SALTLENS DSS: AN EVIDENCE-GATED DIGITAL TWIN FOR "
        "CLIMATE-RESILIENT SALT-PAN MANAGEMENT” is the bonafide work of “SACHIN S B (61782323110095), "
        "THRISHNASANJAY S (61782323110117)” who carried out the project work U23IT705 – PROJECT WORK "
        "under my supervision."
    )
    r_body.font.name = 'Times New Roman'
    r_body.font.size = Pt(14)
    r_body.font.color.rgb = COLOR_BLACK

    # Supervisor & HOD Signatures Block
    # Appendix 2 rule: The term ‘SUPERVISOR’ must be typed in capital letters between the supervisor’s name and academic designation.
    tbl_b_sig = doc.add_table(rows=1, cols=2)
    tbl_b_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_b_sig, color="FFFFFF", sz="0", val="none")
    prevent_row_split(tbl_b_sig.rows[0])
    c_s0, c_s1 = tbl_b_sig.rows[0].cells
    c_s0.width = Inches(3.0)
    c_s1.width = Inches(3.0)

    p_s0 = c_s0.paragraphs[0]
    p_s0.paragraph_format.line_spacing = 1.3
    r = p_s0.add_run(
        "Signature of the Supervisor\n\n\n\n"
        "Mr. B. Mohanraj, M.E., (Ph.D.)\n"
        "SUPERVISOR\n"
        "Assistant Professor\n"
        "Department of Information Technology\n"
        "Sona College of Technology,\n"
        "Salem – 636005"
    )
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11.5)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    p_s1 = c_s1.paragraphs[0]
    p_s1.paragraph_format.line_spacing = 1.3
    p_s1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p_s1.add_run(
        "Head of the Department\n\n\n\n"
        "Dr. J. Akilandeswari, M.E., Ph.D.\n"
        "HEAD OF THE DEPARTMENT\n"
        "Professor and Head\n"
        "Department of Information Technology\n"
        "Sona College of Technology,\n"
        "Salem – 636005"
    )
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11.5)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    p_viva = doc.add_paragraph()
    p_viva.paragraph_format.space_before = Pt(36)
    p_viva.paragraph_format.space_after = Pt(28)
    p_viva.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p_viva.add_run("Submitted for the Project Viva-Voce Examination held on ____________________")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    tbl_exam = doc.add_table(rows=1, cols=2)
    tbl_exam.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_exam, color="FFFFFF", sz="0", val="none")
    prevent_row_split(tbl_exam.rows[0])
    ce0, ce1 = tbl_exam.rows[0].cells
    ce0.width = Inches(3.0)
    ce1.width = Inches(3.0)

    pe0 = ce0.paragraphs[0]
    pe0.paragraph_format.space_before = Pt(12)
    r = pe0.add_run("INTERNAL EXAMINER")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    pe1 = ce1.paragraphs[0]
    pe1.paragraph_format.space_before = Pt(12)
    pe1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = pe1.add_run("EXTERNAL EXAMINER")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    # ==========================================
    # SECTION 1: FRONT MATTER (lowerRoman pagination i, ii, iii...)
    # ==========================================
    sec1 = doc.add_section()
    sec1.page_width = Inches(8.27)
    sec1.page_height = Inches(11.69)
    sec1.top_margin = Inches(1.0)
    sec1.bottom_margin = Inches(1.0)
    sec1.left_margin = Inches(1.25)
    sec1.right_margin = Inches(1.0)
    
    sectPr1 = sec1._sectPr
    pgNumType1 = parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="1" w:fmt="lowerRoman"/>')
    sectPr1.append(pgNumType1)
    add_page_number_to_footer(sec1)

    # 4. DECLARATION FORM
    print("Generating Declaration Form...")
    tbl_dec_hdr = doc.add_table(rows=1, cols=3)
    tbl_dec_hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_dec_hdr, color="FFFFFF", sz="0", val="none")
    prevent_row_split(tbl_dec_hdr.rows[0])
    cd_l, cd_m, cd_r = tbl_dec_hdr.rows[0].cells
    cd_l.width = Inches(1.0)
    cd_m.width = Inches(4.0)
    cd_r.width = Inches(1.0)

    if os.path.exists("extracted_header_logos/col_0_image1.jpeg"):
        cd_l.paragraphs[0].add_run().add_picture("extracted_header_logos/col_0_image1.jpeg", width=Inches(0.95))
    if os.path.exists("extracted_header_logos/col_2_image2.jpeg"):
        cd_r.paragraphs[0].add_run().add_picture("extracted_header_logos/col_2_image2.jpeg", width=Inches(1.05))

    pd_m = cd_m.paragraphs[0]
    pd_m.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = pd_m.add_run("Sona College of Technology\nDepartment of Information Technology\nU23IT705/U23ADS705 - PROJECT - DECLARATION FORM")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    p_d_sub = doc.add_paragraph()
    p_d_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_d_sub.paragraph_format.space_before = Pt(16)
    p_d_sub.paragraph_format.space_after = Pt(16)
    r = p_d_sub.add_run("Declaration by the Student(s)")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    add_body_p(
        "We hereby declare that the project titled “SALTLENS DSS: AN EVIDENCE-GATED DIGITAL TWIN "
        "FOR CLIMATE-RESILIENT SALT-PAN MANAGEMENT” is an original work carried out by us under the "
        "guidance of Mr. B. Mohanraj, Assistant Professor, Department of Information Technology, "
        "Sona College of Technology (Autonomous), Salem. This project has not been submitted elsewhere "
        "for the award of any degree or diploma.",
        space_after=10, line_spacing=1.5
    )

    add_body_p(
        "We also affirm that all the data and information used in this project are genuine and properly "
        "acknowledged. The field records referred to in this report were reported by the site operator "
        "of the North Pan, Salem, and are presented with the provenance, evidence categories, and limitations "
        "recorded in the system. This project work is carried out with reference to the United Nations Sustainable "
        "Development Goals: SDG 13 – Climate Action, SDG 9 – Industry, Innovation, and Infrastructure, "
        "SDG 12 – Responsible Consumption and Production, and SDG 8 – Decent Work and Economic Growth.",
        space_after=20, line_spacing=1.5
    )

    p_dp = doc.add_paragraph()
    p_dp.paragraph_format.space_after = Pt(14)
    r = p_dp.add_run("Date: October 2026\nPlace: Sona College of Technology, Salem")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    # Students Sign Table
    tbl_ds = doc.add_table(rows=3, cols=4)
    tbl_ds.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_ds, color=HEX_BORDER, sz="4", val="single")
    d_hdrs = ["S.No", "Register Number", "Name", "Signature"]
    d_widths = [Inches(0.6), Inches(1.8), Inches(2.2), Inches(1.4)]
    
    for i, title in enumerate(d_hdrs):
        cell = tbl_ds.rows[0].cells[i]
        cell.width = d_widths[i]
        set_cell_background(cell, HEX_HDR_BG)
        set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10.5)
        r.bold = True
        r.font.color.rgb = COLOR_BLACK

    d_data = [
        ("1.", "61782323110095", "SACHIN S B", ""),
        ("2.", "61782323110117", "THRISHNASANJAY S", "")
    ]
    for r_idx, row in enumerate(d_data):
        r_el = tbl_ds.rows[r_idx + 1]
        prevent_row_split(r_el)
        for c_idx, val in enumerate(row):
            cell = r_el.cells[c_idx]
            cell.width = d_widths[c_idx]
            set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
            p = cell.paragraphs[0]
            if c_idx in [0, 1]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10)
            r.font.color.rgb = COLOR_BLACK

    doc.add_page_break()

    # 5. ACKNOWLEDGMENT
    print("Generating Acknowledgment...")
    p_ack_hdr = doc.add_paragraph()
    p_ack_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ack_hdr.paragraph_format.space_before = Pt(14)
    p_ack_hdr.paragraph_format.space_after = Pt(16)
    r = p_ack_hdr.add_run("ACKNOWLEDGMENT")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    add_body_p(
        "We express our sincere gratitude to the Management of Sona College of Technology (Autonomous), Salem, "
        "for providing the state-of-the-art infrastructure, computing resources, and academic environment in "
        "which this project was carried out. We express our heartfelt gratitude to our respected Principal for "
        "the continuous encouragement, institutional leadership, and supportive environment extended to student "
        "innovations.", line_spacing=1.5
    )
    add_body_p(
        "We are profoundly grateful to Dr. J. Akilandeswari, Professor and Head of the Department of Information "
        "Technology, for providing excellent laboratory facilities, administrative support, and scholastic guidance "
        "that enabled this project work to proceed smoothly to completion.", line_spacing=1.5
    )
    add_body_p(
        "We place on record our deep sense of gratitude and respect to our project supervisor, Mr. B. Mohanraj, "
        "Assistant Professor, Department of Information Technology, for his patient direction, insightful technical "
        "critique of every draft, and steadfast insistence that every engineering claim in this report be tied directly "
        "to the empirical and computational evidence that supports it. That methodological insistence formed the core "
        "architectural foundation of the SaltLens decision support system.", line_spacing=1.5
    )
    add_body_p(
        "We extend our sincere thanks to the site operator of the North Pan, Salem, whose notebook records and "
        "reported field readings made the empirical case study possible, and who permitted those readings to be "
        "presented in this report exactly as reported, with their physical and operational limitations explicitly recorded.",
        line_spacing=1.5
    )
    add_body_p(
        "We also convey our sincere thanks to all faculty members of the Department of Information Technology, "
        "our project review panel, our fellow classmates, and our family members for their constant encouragement "
        "and moral support throughout the course of this work.", line_spacing=1.5
    )

    doc.add_page_break()

    # 6. ABSTRACT
    print("Generating Abstract...")
    p_abs_hdr = doc.add_paragraph()
    p_abs_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_abs_hdr.paragraph_format.space_before = Pt(14)
    p_abs_hdr.paragraph_format.space_after = Pt(16)
    r = p_abs_hdr.add_run("ABSTRACT")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    # The Abstract should be one page synopsis typed double line spacing / 1.5 spacing, Font Size 14 / 12
    add_body_p(
        "Rainfall can rapidly dilute hypersaline brine in solar salt pans, but a weather forecast alone cannot "
        "establish how a managed pond will respond. This project presents SaltLens DSS, an implemented, "
        "evidence-gated decision-support prototype for salt-pan management. The system combines timestamped pan "
        "observations, weather records, a deterministic rainfall–dilution mass balance, versioned operator-configured "
        "rules, recorded operator decisions and an auditable field-validation workflow. It is built with a Next.js user "
        "interface, FastAPI services and MongoDB Atlas storage.", line_spacing=1.5
    )
    add_body_p(
        "The central design principle is an explicit boundary between what was observed, what was forecast, what a "
        "simplified physical model calculated, what a rule recommended, and what an operator decided. Each stage "
        "carries its own evidentiary status and none silently replaces another. A machine-learning evidence gate blocks "
        "predictive output until a genuine labelled dataset and an independently evaluated model exist; consequently "
        "no machine-learning accuracy figure is reported.", line_spacing=1.5
    )
    add_body_p(
        "A North Pan case in Salem, India, uses a notebook-sourced pre-event record of 250 g/L salinity and 20 cm depth "
        "on 13 March 2026 at 13:30 IST. The reported rainfall event began on 14 March at 23:30 IST, with 90 mm of "
        "rainfall reported. The rainfall-only replay predicts 172.41 g/L and 29 cm, whereas operator-reported post-event "
        "readings are 20 g/L and 50 cm. An outbound brine transfer of unknown volume occurred and the post-event readings "
        "lack independent corroboration, so the case is descriptive and cannot establish model accuracy. A separate "
        "September 2026, 25 mm forecast scenario predicts 222.22 g/L and demonstrates the versioned rule-based "
        "recommendation and a recorded operator acceptance, without a documented physical harvest.", line_spacing=1.5
    )
    add_body_p(
        "The prototype therefore demonstrates reproducible software behaviour, provenance control and audit-preserving "
        "reconciliation. It does not demonstrate generalised physical prediction, trained machine-learning performance "
        "or operational benefit; each of these requires separate empirical evidence, and the report specifies the studies "
        "needed to obtain it.", line_spacing=1.5
    )

    p_kw = doc.add_paragraph()
    p_kw.paragraph_format.space_before = Pt(8)
    p_kw.paragraph_format.line_spacing = 1.3
    r1 = p_kw.add_run("Keywords: ")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(12)
    r1.bold = True
    r1.font.color.rgb = COLOR_BLACK
    r2 = p_kw.add_run("solar salt pans, digital twin, decision support, provenance, rainfall dilution, field reporting, audit trail")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(12)
    r2.font.color.rgb = COLOR_BLACK

    doc.add_page_break()

    # 7. TABLE OF CONTENTS
    # Appendix 3 rule: "The title page and Bonafide Certificate will not find a place among the items listed in the Table of Contents but the page numbers of which are in lower case Roman letters."
    print("Generating Table of Contents...")
    p_toc_hdr = doc.add_paragraph()
    p_toc_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_toc_hdr.paragraph_format.space_before = Pt(14)
    p_toc_hdr.paragraph_format.space_after = Pt(12)
    r = p_toc_hdr.add_run("TABLE OF CONTENTS")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    add_dot_leader_header("TITLE", "PAGE NO.")

    toc_entries = [
        ("DECLARATION", "i", True, 0),
        ("ACKNOWLEDGMENT", "ii", True, 0),
        ("ABSTRACT", "iii", True, 0),
        ("LIST OF TABLES", "vi", True, 0),
        ("LIST OF FIGURES", "vii", True, 0),
        ("LIST OF ABBREVIATIONS", "viii", True, 0),
        ("1. INTRODUCTION", "1", True, 0),
        ("1.1 Background of the Project", "1", False, 1),
        ("1.2 Problem Statement", "2", False, 1),
        ("1.3 Motivation", "3", False, 1),
        ("1.4 Objectives", "3", False, 1),
        ("1.5 Scope of the Project", "4", False, 1),
        ("2. LITERATURE REVIEW", "6", True, 0),
        ("2.1 Introduction", "6", False, 1),
        ("2.2 Review of Related Studies", "6", False, 1),
        ("2.3 Techniques and Technologies Used in Existing Systems", "10", False, 1),
        ("2.4 Research Gap", "11", False, 1),
        ("2.5 Summary", "12", False, 1),
        ("3. CHAPTER 1 – NEED FOR PROPOSED TECHNIQUE", "13", True, 0),
        ("3.1 Existing Techniques", "13", False, 1),
        ("3.2 Limitations of Existing Techniques", "15", False, 1),
        ("3.3 Proposed Technique", "16", False, 1),
        ("3.4 Objectives of the Proposed Technique", "17", False, 1),
        ("3.5 Advantages of the Proposed Technique", "17", False, 1),
        ("4. CHAPTER 2 – DESIGN OF PROPOSED SYSTEM AND ITS IMPLEMENTATION", "18", True, 0),
        ("4.1 Introduction", "18", False, 1),
        ("4.2 System Requirements", "18", False, 1),
        ("4.2.1 Hardware Requirements", "18", False, 2),
        ("4.2.2 Software Requirements", "19", False, 2),
        ("4.3 System Architecture", "20", False, 1),
        ("4.4 System Components / Modules", "21", False, 1),
        ("4.4.1 Module 1: Pan-Specific Digital Twin Module", "21", False, 2),
        ("4.4.2 Module 2: Deterministic Rainfall-Dilution Engine", "22", False, 2),
        ("4.4.3 Module 3: Versioned Rule Engine", "22", False, 2),
        ("4.4.4 Module 4: Machine-Learning Evidence Gate", "23", False, 2),
        ("4.4.5 Module 5: Field-Event Reconciliation and Audit Engine", "23", False, 2),
        ("4.5 Proposed Methodology", "24", False, 1),
        ("4.6 System Workflow", "24", False, 1),
        ("4.7 Algorithms / Techniques Used", "25", False, 1),
        ("4.8 Dataset / Data Collection", "27", False, 1),
        ("4.9 Data Preprocessing", "28", False, 1),
        ("4.10 Model / System Development", "29", False, 1),
        ("4.11 System Design Diagrams", "30", False, 1),
        ("4.12 Implementation", "31", False, 1),
        ("4.13 Coding / Program Implementation", "32", False, 1),
        ("4.14 User Interface / Application Screens", "36", False, 1),
        ("4.15 Chapter Summary", "37", False, 1),
        ("5. CHAPTER 3 – RESULTS AND DISCUSSION", "38", True, 0),
        ("5.1 Introduction", "38", False, 1),
        ("5.2 Experimental Setup", "38", False, 1),
        ("5.3 Evaluation Metrics", "39", False, 1),
        ("5.4 Experimental Results", "40", False, 1),
        ("5.5 Comparison with Existing Techniques", "44", False, 1),
        ("5.6 Discussion of Results", "46", False, 1),
        ("5.7 Chapter Summary", "48", False, 1),
        ("6. CHAPTER 4 – CONCLUSION AND FUTURE SCOPE", "49", True, 0),
        ("6.1 Summary of the Project", "49", False, 1),
        ("6.2 Conclusion", "50", False, 1),
        ("6.3 Future Scope", "51", False, 1),
        ("SDG GOALS", "53", True, 0),
        ("PUBLICATIONS (IF ANY)", "55", True, 0),
        ("REFERENCES", "56", True, 0),
        ("APPENDIX 1: SENSITIVITY TABLES", "58", True, 0),
        ("APPENDIX 2: HAND-VERIFICATION WORKSHEET", "60", True, 0),
        ("APPENDIX 3: FIELD DATA-CAPTURE TEMPLATE", "61", True, 0),
        ("APPENDIX 4: GLOSSARY AND ABBREVIATIONS", "62", True, 0),
        ("APPENDIX 5: NOTATION", "63", True, 0)
    ]
    for title, pg, is_b, ind in toc_entries:
        add_dot_leader_line(title, pg, is_bold=is_b, indent_spaces=ind)

    doc.add_page_break()

    # 8. LIST OF TABLES
    # Template rule: "The list should use exactly the same captions as they appear above the tables in the text. One and a half spacing should be adopted for typing the matter under this head."
    print("Generating List of Tables...")
    p_lot_hdr = doc.add_paragraph()
    p_lot_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_lot_hdr.paragraph_format.space_before = Pt(14)
    p_lot_hdr.paragraph_format.space_after = Pt(12)
    r = p_lot_hdr.add_run("LIST OF TABLES")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    add_dot_leader_header("TABLE NO. / TITLE", "PAGE NO.")

    lot_entries = [
        ("Table 2.1  Comparison of related research and SaltLens DSS", "10"),
        ("Table 3.1  Summary of existing techniques for salt-pan rainfall management", "14"),
        ("Table 3.2  Limitations of existing techniques and the corresponding SaltLens mechanism", "16"),
        ("Table 4.1  Functional requirements of SaltLens DSS", "19"),
        ("Table 4.2  Non-functional requirements of SaltLens DSS", "19"),
        ("Table 4.3  Technology stack and system roles", "20"),
        ("Table 4.4  Evidence categories used by the interface", "21"),
        ("Table 4.5  Principal record types and provenance role", "22"),
        ("Table 5.1  Separate deterministic rainfall-only scenarios", "40"),
        ("Table 5.2  March North Pan event record (as reported)", "41"),
        ("Table 5.3  Implied water volume and dissolved-salt mass for the March event (12,300 m² pan)", "43"),
        ("Table 5.4  What software verification supports and does not support", "44"),
        ("Table 5.5  Qualitative comparison of approaches by design property", "45"),
        ("Table 5.6  Information to retain for each event in a multi-event validation study", "47"),
        ("Table 5.7  Claim types and the evidence required to support them", "48"),
        ("Table A.1  Predicted post-rain salinity C₁ (g/L) for C₀ = 250 g/L", "58"),
        ("Table A.2  Predicted percentage dilution δ (%) (independent of C₀)", "58"),
        ("Table A.3  Predicted post-rain depth D₁ (cm)", "59"),
        ("Table B.1  Worked verification for the two scenarios (A = 12,300 m², C₀ = 250 g/L, D₀ = 20 cm)", "60"),
        ("Table C.1  Per-event field data-capture template", "61"),
        ("Table D.1  Glossary and abbreviations", "62"),
        ("Table E.1  Symbols used in the mass-balance equations", "63")
    ]
    for title, pg in lot_entries:
        add_dot_leader_line(title, pg, is_bold=False, indent_spaces=0)

    doc.add_page_break()

    # 9. LIST OF FIGURES
    # Template rule: "The list should use exactly the same captions as they appear below the figures in the text. One and a half spacing should be adopted for typing the matter under this head."
    print("Generating List of Figures...")
    p_lof_hdr = doc.add_paragraph()
    p_lof_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_lof_hdr.paragraph_format.space_before = Pt(14)
    p_lof_hdr.paragraph_format.space_after = Pt(12)
    r = p_lof_hdr.add_run("LIST OF FIGURES")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    add_dot_leader_header("FIGURE NO. / TITLE", "PAGE NO.")

    lof_entries = [
        ("Figure 4.1  Layered architecture of SaltLens DSS", "20"),
        ("Figure 4.2  Evidence chain: each stage carries its own evidentiary status", "22"),
        ("Figure 4.3  North Pan dashboard with pan-specific recorded state and evidence labels", "24"),
        ("Figure 4.4  Predicted salinity versus rainfall for several initial depths (C₀ = 250 g/L)", "26"),
        ("Figure 4.5  Percentage dilution versus rainfall; independent of initial salinity", "27"),
        ("Figure 4.6  Rule evaluation flow for the harvest/flood rule", "28"),
        ("Figure 4.7  Versioned rule recommendation; interface reports no trained ML estimate was used", "29"),
        ("Figure 4.8  Operator acceptance recorded for workflow testing; no physical harvest occurred", "30"),
        ("Figure 4.9  September 25 mm forecast simulation showing 222.22 g/L predicted salinity", "36"),
        ("Figure 5.1  Status of the validity dimensions in the present study", "39"),
        ("Figure 5.2  March North Pan salinity: rainfall-only prediction of 172.41 g/L versus reported 20 g/L", "42"),
        ("Figure 5.3  March North Pan depth: rainfall-only prediction of 29 cm versus reported 50 cm", "42"),
        ("Figure 5.4  March event: rainfall-only prediction compared with operator-reported values", "43"),
        ("Figure 5.5  Implied water volume and dissolved-salt mass for the March event", "44")
    ]
    for title, pg in lof_entries:
        add_dot_leader_line(title, pg, is_bold=False, indent_spaces=0)

    doc.add_page_break()

    # 10. LIST OF ABBREVIATIONS
    print("Generating List of Abbreviations...")
    p_loa_hdr = doc.add_paragraph()
    p_loa_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_loa_hdr.paragraph_format.space_before = Pt(14)
    p_loa_hdr.paragraph_format.space_after = Pt(14)
    r = p_loa_hdr.add_run("LIST OF ABBREVIATIONS")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    abbr_items = [
        ("AI", "Artificial Intelligence"),
        ("API", "Application Programming Interface"),
        ("CORS", "Cross-Origin Resource Sharing"),
        ("CSV", "Comma-Separated Values"),
        ("DFD", "Data Flow Diagram"),
        ("DSS", "Decision Support System"),
        ("ER", "Entity-Relationship"),
        ("FASTAPI", "Modern, Fast Web Framework for Building APIs with Python"),
        ("IoT", "Internet of Things"),
        ("IST", "Indian Standard Time (UTC+05:30)"),
        ("JSON", "JavaScript Object Notation"),
        ("LBPH", "Local Binary Patterns Histograms"),
        ("ML", "Machine Learning"),
        ("REST", "Representational State Transfer"),
        ("RFID", "Radio-Frequency Identification"),
        ("SDG", "Sustainable Development Goals (United Nations)"),
        ("UI", "User Interface"),
        ("UTC", "Coordinated Universal Time"),
        ("XAI", "Explainable Artificial Intelligence")
    ]

    tbl_abbr = doc.add_table(rows=len(abbr_items) + 1, cols=2)
    tbl_abbr.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_abbr, color="FFFFFF", sz="0", val="none")
    tbl_abbr.rows[0].cells[0].width = Inches(2.0)
    tbl_abbr.rows[0].cells[1].width = Inches(4.0)
    
    p_ah0 = tbl_abbr.rows[0].cells[0].paragraphs[0]
    r = p_ah0.add_run("ABBREVIATION")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    p_ah1 = tbl_abbr.rows[0].cells[1].paragraphs[0]
    r = p_ah1.add_run("EXPANSION")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r.bold = True
    r.font.color.rgb = COLOR_BLACK

    for idx, (abbr, exp) in enumerate(abbr_items):
        row = tbl_abbr.rows[idx + 1]
        prevent_row_split(row)
        c0, c1 = row.cells
        c0.width = Inches(2.0)
        c1.width = Inches(4.0)
        set_cell_margins(c0, top=30, bottom=30, left=50, right=50)
        set_cell_margins(c1, top=30, bottom=30, left=50, right=50)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.line_spacing = 1.3
        r0 = p0.add_run(abbr)
        r0.font.name = 'Times New Roman'
        r0.font.size = Pt(11)
        r0.bold = True
        r0.font.color.rgb = COLOR_BLACK

        p1 = c1.paragraphs[0]
        p1.paragraph_format.line_spacing = 1.3
        r1 = p1.add_run(exp)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_BLACK

    # ==========================================
    # SECTION 2: MAIN TEXT (decimal pagination 1, 2, 3...)
    # ==========================================
    sec2 = doc.add_section()
    sec2.page_width = Inches(8.27)
    sec2.page_height = Inches(11.69)
    sec2.top_margin = Inches(1.0)
    sec2.bottom_margin = Inches(1.0)
    sec2.left_margin = Inches(1.25)
    sec2.right_margin = Inches(1.0)
    
    sectPr2 = sec2._sectPr
    pgNumType2 = parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="1" w:fmt="decimal"/>')
    sectPr2.append(pgNumType2)
    add_page_number_to_footer(sec2)

    return doc, {
        'all_tables_json': all_tables_json,
        'add_chapter_heading': add_chapter_heading,
        'add_sec_heading_2': add_sec_heading_2,
        'add_sec_heading_3': add_sec_heading_3,
        'add_sec_heading_4': add_sec_heading_4,
        'add_body_p': add_body_p,
        'add_bullet_p': add_bullet_p,
        'add_figure_item': add_figure_item,
        'add_normal_table': add_normal_table,
        'set_table_borders': set_table_borders,
        'set_cell_background': set_cell_background,
        'set_cell_margins': set_cell_margins,
        'prevent_row_split': prevent_row_split,
        'set_repeat_header': set_repeat_header,
        'HEX_BORDER': HEX_BORDER,
        'HEX_HDR_BG': HEX_HDR_BG
    }

if __name__ == '__main__':
    doc, helpers = build_full_strict_report()
    doc.save("test_front_matter.docx")
    print("Front matter generated successfully in test_front_matter.docx")
