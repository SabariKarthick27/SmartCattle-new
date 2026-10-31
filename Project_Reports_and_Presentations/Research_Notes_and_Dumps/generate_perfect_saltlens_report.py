import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import json
import win32com.client

def generate_report():
    print("=" * 70)
    print("Starting Strict SaltLens DSS Project Report Generation")
    print("Strictly adhering to SAMPLE PROJECT REPORT FORMAT & INSTRUCTIONS.docx")
    print("=" * 70)

    doc = docx.Document()

    TAB_STOP_RIGHT = 8640  # 6.0 in right tab stop for dot leaders
    TAB_STOP_ABBR = 2160   # 1.5 in tab stop

    # Configure Section 0 (Cover & Title Page - No headers/footers, no page numbers)
    sec0 = doc.sections[0]
    sec0.page_width = Inches(8.27)
    sec0.page_height = Inches(11.69)
    sec0.top_margin = Inches(1.0)
    sec0.bottom_margin = Inches(1.0)
    sec0.left_margin = Inches(1.25)
    sec0.right_margin = Inches(1.0)

    COLOR_BLACK = RGBColor(0x00, 0x00, 0x00)
    COLOR_DARK = RGBColor(0x1A, 0x1A, 0x1A)

    HEX_HDR_BG = "EFEFEF"
    HEX_ALT_BG = "FAFAFA"
    HEX_BORDER = "888888"

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = COLOR_BLACK

    # XML Helper Functions
    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=70, bottom=70, left=100, right=100):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def prevent_row_split(row):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    def set_repeat_header(row):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    def set_table_borders(table, color=HEX_BORDER, sz="4", val="single"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def add_page_number_to_footer(section):
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.text = ""
        r = p.add_run()
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        r.font.color.rgb = COLOR_BLACK
        fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
        instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
        fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
        fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
        r._r.append(fldChar1)
        r._r.append(instrText)
        r._r.append(fldChar2)
        r._r.append(fldChar3)

    # Paragraph helpers
    def add_chapter_heading(title_text, chapter_number=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        if chapter_number is not None:
            r1 = p.add_run(f"CHAPTER {chapter_number}\n")
            r1.font.name = 'Times New Roman'
            r1.font.size = Pt(15)
            r1.bold = True
            r1.font.color.rgb = COLOR_BLACK
            
        r2 = p.add_run(title_text.upper())
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(14)
        r2.bold = True
        r2.font.color.rgb = COLOR_BLACK
        return p

    def add_sec_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13)
        r.bold = True
        r.font.color.rgb = COLOR_BLACK
        return p

    def add_sec_heading_3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.bold = True
        r.font.color.rgb = COLOR_BLACK
        return p

    def add_sec_heading_4(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11.5)
        r.bold = True
        r.font.color.rgb = COLOR_BLACK
        return p

    def add_body_p(text, bold_prefix="", space_after=6, space_before=0, line_spacing=1.5, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = 'Times New Roman'
            r_pre.font.size = Pt(12)
            r_pre.bold = True
            r_pre.font.color.rgb = COLOR_BLACK
            
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.italic = italic
        r.font.color.rgb = COLOR_BLACK
        return p

    def add_list_item(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.35
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = 'Times New Roman'
            r_pre.font.size = Pt(12)
            r_pre.bold = True
            r_pre.font.color.rgb = COLOR_BLACK
            
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.font.color.rgb = COLOR_BLACK
        return p

    def add_figure_item(image_path, caption_text, width_inches=5.2):
        if os.path.exists(image_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(10)
            p_img.paragraph_format.space_after = Pt(4)
            p_img.paragraph_format.keep_with_next = True
            run_img = p_img.add_run()
            run_img.add_picture(image_path, width=Inches(width_inches))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(12)
            r_cap = p_cap.add_run(caption_text)
            r_cap.font.name = 'Times New Roman'
            r_cap.font.size = Pt(11)
            r_cap.bold = True
            r_cap.font.color.rgb = COLOR_BLACK

    def add_normal_table(headers, data, col_widths, caption_text=""):
        if caption_text:
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(12)
            p_cap.paragraph_format.space_after = Pt(4)
            p_cap.paragraph_format.keep_with_next = True
            r_cap = p_cap.add_run(caption_text)
            r_cap.font.name = 'Times New Roman'
            r_cap.font.size = Pt(11)
            r_cap.bold = True
            r_cap.font.color.rgb = COLOR_BLACK

        table = doc.add_table(rows=len(data) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, color=HEX_BORDER, sz="4", val="single")

        hdr_row = table.rows[0]
        set_repeat_header(hdr_row)
        prevent_row_split(hdr_row)

        for i, title in enumerate(headers):
            cell = hdr_row.cells[i]
            cell.width = col_widths[i]
            set_cell_background(cell, HEX_HDR_BG)
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(title)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10)
            r.bold = True
            r.font.color.rgb = COLOR_BLACK

        for r_idx, row_data in enumerate(data):
            row = table.rows[r_idx + 1]
            prevent_row_split(row)
            bg_color = HEX_ALT_BG if r_idx % 2 == 1 else "FFFFFF"

            for c_idx, val in enumerate(row_data):
                cell = row.cells[c_idx]
                cell.width = col_widths[c_idx]
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.2
                
                # Center numeric / short items
                if c_idx == 0 and len(headers) > 2 and len(val) < 20:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif "%" in val or val.replace(".", "").isdigit() or "mm" in val or "cm" in val or "g/L" in val:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

                r = p.add_run(val)
                r.font.name = 'Times New Roman'
                r.font.size = Pt(9.5)
                r.font.color.rgb = COLOR_BLACK

        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_before = Pt(1)
        p_after.paragraph_format.space_after = Pt(8)

    def add_dot_leader_line(title_text, page_str, is_bold=False, indent_spaces=0):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.3
        pPr = p._p.get_or_add_pPr()
        tabs = parse_xml(f'<w:tabs {nsdecls("w")}><w:tab w:val="right" w:leader="dot" w:pos="{TAB_STOP_RIGHT}"/></w:tabs>')
        pPr.append(tabs)

        prefix = "    " * indent_spaces
        r1 = p.add_run(prefix + title_text)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(11)
        if is_bold:
            r1.bold = True
        r1.font.color.rgb = COLOR_BLACK

        r2 = p.add_run(f"\t{page_str}")
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(11)
        if is_bold:
            r2.bold = True
        r2.font.color.rgb = COLOR_BLACK
        return p

    def add_dot_leader_header(left_title="TITLE", right_title="PAGE NO."):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        pPr = p._p.get_or_add_pPr()
        tabs = parse_xml(f'<w:tabs {nsdecls("w")}><w:tab w:val="right" w:pos="{TAB_STOP_RIGHT}"/></w:tabs>')
        pPr.append(tabs)

        r1 = p.add_run(left_title)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(11.5)
        r1.bold = True
        r1.font.color.rgb = COLOR_BLACK

        r2 = p.add_run(f"\t{right_title}")
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(11.5)
        r2.bold = True
        r2.font.color.rgb = COLOR_BLACK

    # Load extracted tables from JSON
    with open('saltlens_tables_extracted.json', 'r', encoding='utf-8') as f:
        all_tables_json = json.load(f)

    print("Loaded table datasets.")
    return doc, {
        'sec0': sec0,
        'set_cell_background': set_cell_background,
        'set_cell_margins': set_cell_margins,
        'prevent_row_split': prevent_row_split,
        'set_repeat_header': set_repeat_header,
        'set_table_borders': set_table_borders,
        'add_page_number_to_footer': add_page_number_to_footer,
        'add_chapter_heading': add_chapter_heading,
        'add_sec_heading_2': add_sec_heading_2,
        'add_sec_heading_3': add_sec_heading_3,
        'add_sec_heading_4': add_sec_heading_4,
        'add_body_p': add_body_p,
        'add_list_item': add_list_item,
        'add_figure_item': add_figure_item,
        'add_normal_table': add_normal_table,
        'add_dot_leader_line': add_dot_leader_line,
        'add_dot_leader_header': add_dot_leader_header,
        'all_tables_json': all_tables_json,
        'COLOR_BLACK': COLOR_BLACK,
        'TAB_STOP_RIGHT': TAB_STOP_RIGHT,
        'TAB_STOP_ABBR': TAB_STOP_ABBR
    }

if __name__ == '__main__':
    generate_report()
