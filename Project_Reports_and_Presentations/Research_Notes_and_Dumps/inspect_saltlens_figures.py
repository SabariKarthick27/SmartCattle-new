import docx

doc = docx.Document('SaltLens_DSS_Final_Project_Report_Final_With_Cover.docx')
print('=== Paragraphs with drawings / images ===')
for i, p in enumerate(doc.paragraphs):
    has_img = any('graphic' in r._r.xml for r in p.runs)
    if has_img:
        prev_p = doc.paragraphs[i-1].text.strip() if i > 0 else ''
        next_p = doc.paragraphs[i+1].text.strip() if i+1 < len(doc.paragraphs) else ''
        next2_p = doc.paragraphs[i+2].text.strip() if i+2 < len(doc.paragraphs) else ''
        print(f'P{i}: prev="{prev_p[:50]}", next="{next_p[:60]}", next2="{next2_p[:60]}"')
