import docx

doc = docx.Document(r'Project_Report_Final_20_10_26 (1).docx')
for i in range(200, 241):
    if i < len(doc.paragraphs):
        p = doc.paragraphs[i]
        has_draw = any('drawing' in r._element.xml for r in p.runs)
        txt = p.text.encode('ascii', errors='replace').decode('ascii')
        print(f"P{i} [draw={has_draw}]: {txt}")
