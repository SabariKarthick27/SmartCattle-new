import docx
import os
import re

doc = docx.Document('mini project report.docx')
t0 = doc.tables[0]
os.makedirs('extracted_header_logos', exist_ok=True)

for i in range(3):
    cell = t0.rows[0].cells[i]
    print(f"Col {i} width: {cell.width}")
    for p_idx, p in enumerate(cell.paragraphs):
        print(f"  P{p_idx} text: '{p.text}'")
        for r_idx, r in enumerate(p.runs):
            xml = r._r.xml
            if 'blip' in xml:
                embed_ids = re.findall(r'embed="([^"]+)"', xml)
                for eid in embed_ids:
                    rel = doc.part.rels[eid]
                    target_ref = rel.target_ref
                    img_blob = rel.target_part.blob
                    out_name = f"col_{i}_{os.path.basename(target_ref)}"
                    out_path = os.path.join('extracted_header_logos', out_name)
                    with open(out_path, 'wb') as f:
                        f.write(img_blob)
                    print(f"    Saved logo from Col {i}: {out_path} ({len(img_blob)} bytes)")
