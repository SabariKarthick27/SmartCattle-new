import docx
import sys
import os
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r'c:\Project\SmartCattle new\QR MINI PROJECT REPORT.docx')

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

# Dump all paragraphs with text to a file for complete analysis
with open('current_report_dump.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt:
            f.write(f"[{i}] {txt}\n")

print("Dumped current report to current_report_dump.txt")
