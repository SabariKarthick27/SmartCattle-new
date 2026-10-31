import docx
import sys
import os
import io
from PIL import Image
sys.stdout.reconfigure(encoding='utf-8')

doc_path = r'c:\Project\SmartCattle new\QR MINI PROJECT REPORT.docx'
doc = docx.Document(doc_path)

extract_dir = r'c:\Project\SmartCattle new\report_extracted_figures'
os.makedirs(extract_dir, exist_ok=True)

print("="*80)
print(f"EXHAUSTIVE AUDIT OF ALL FIGURES & DRAWINGS IN: {doc_path}")
print("="*80)

# Extract images from docx package parts
image_part_map = {}
for rel_id, rel in doc.part.rels.items():
    if "image" in rel.target_ref:
        img_part = rel.target_part
        img_bytes = img_part.blob
        filename = os.path.basename(rel.target_ref)
        filepath = os.path.join(extract_dir, filename)
        with open(filepath, 'wb') as f:
            f.write(img_bytes)
        
        # Get image dimensions using PIL
        try:
            im = Image.open(io.BytesIO(img_bytes))
            dimensions = f"{im.width}x{im.height} ({im.format})"
        except Exception as e:
            dimensions = f"Unknown ({e})"
            
        image_part_map[rel_id] = {
            'filename': filename,
            'size_bytes': len(img_bytes),
            'dimensions': dimensions,
            'path': filepath
        }

print(f"Extracted {len(image_part_map)} image assets from docx.")

# Map paragraphs to images
fig_counter = 0
for i, p in enumerate(doc.paragraphs):
    # Find blip elements in run
    blips = p._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
    if blips:
        for blip in blips:
            embed_id = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
            img_info = image_part_map.get(embed_id, {'filename': 'Unknown', 'dimensions': 'N/A', 'size_bytes': 0})
            
            # Look forward/backward for caption
            caption = ""
            for offset in range(-3, 4):
                idx = i + offset
                if 0 <= idx < len(doc.paragraphs):
                    t = doc.paragraphs[idx].text.strip()
                    if t.lower().startswith("fig") or "figure " in t.lower() or "equation" in t.lower():
                        caption = t
                        break
            
            print(f"\nDrawing in P[{i}]:")
            print(f"  Asset: {img_info['filename']} | Size: {img_info['size_bytes']:,} bytes | Resolution: {img_info['dimensions']}")
            print(f"  Paragraph text: '{p.text.strip()[:60]}'")
            print(f"  Associated Caption/Context: '{caption}'")
