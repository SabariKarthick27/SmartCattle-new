import docx

doc = docx.Document(r'Project_Report_Final_20_10_26 (1).docx')
with open('equations_dump.txt', 'w', encoding='utf-8') as out:
    for i, p in enumerate(doc.paragraphs):
        if 'Equation (' in p.text or 'Equation ' in p.text:
            out.write(f"\n==================== P{i} ====================\n")
            out.write(f"HEADER: {p.text}\n")
            for j in range(i+1, min(len(doc.paragraphs), i+8)):
                txt = doc.paragraphs[j].text.strip()
                if txt:
                    out.write(f"  [P{j}] {txt}\n")
print("Dumped equations successfully!")
