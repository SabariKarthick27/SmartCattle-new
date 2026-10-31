import docx
import re

doc = docx.Document(r'c:\Project\SmartCattle new\SmartCattle_Project_Report_PERFECT_MATH.docx')

# Collect text from all paragraphs and all tables except Table 7
texts = []
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        # skip title of abbreviations if paragraph 55
        if 'LIST OF ABBREVIATIONS' in p.text:
            continue
        texts.append(('paragraph', i, p.text))

for t_idx, t in enumerate(doc.tables):
    if t_idx == 7:
        continue # skip abbreviations table
    for r_idx, r in enumerate(t.rows):
        for c_idx, c in enumerate(r.cells):
            if c.text.strip():
                texts.append((f'table_{t_idx}', f'r{r_idx}c{c_idx}', c.text))

abbreviations = [
    ('AI', 'Artificial Intelligence'),
    ('API', 'Application Programming Interface'),
    ('BLE', 'Bluetooth Low Energy'),
    ('BRD', 'Bovine Respiratory Disease'),
    ('CHDI', 'Composite Health Deviation Index'),
    ('CNN', 'Convolutional Neural Network'),
    ('CSS', 'Cascading Style Sheets'),
    ('DFD', 'Data Flow Diagram'),
    ('ER', 'Entity Relationship'),
    ('EWMA', 'Exponentially Weighted Moving Average'),
    ('F1', 'Harmonic Mean of Precision and Recall'),
    ('HTTP', 'Hypertext Transfer Protocol'),
    ('HTTPS', 'Hypertext Transfer Protocol Secure'),
    ('IMU', 'Inertial Measurement Unit'),
    ('IoT', 'Internet of Things'),
    ('JSON', 'JavaScript Object Notation'),
    ('JWT', 'JSON Web Token'),
    ('LF', 'Low Frequency (RFID)'),
    ('ML', 'Machine Learning'),
    ('MQTT', 'Message Queuing Telemetry Transport'),
    ('OCR', 'Optical Character Recognition'),
    ('OLS', 'Ordinary Least Squares'),
    ('ORM', 'Object-Relational Mapping'),
    ('PLF', 'Precision Livestock Farming'),
    ('PWA', 'Progressive Web Application'),
    ('QR', 'Quick Response Code'),
    ('RBAC', 'Role-Based Access Control'),
    ('REST', 'Representational State Transfer'),
    ('RFID', 'Radio-Frequency Identification'),
    ('SDG', 'Sustainable Development Goals'),
    ('SHAP', 'SHapley Additive exPlanations'),
    ('THI', 'Temperature-Humidity Index'),
    ('UHF', 'Ultra-High Frequency (RFID)'),
    ('UI', 'User Interface'),
    ('UN', 'United Nations'),
    ('URL', 'Uniform Resource Locator'),
    ('UUID', 'Universally Unique Identifier'),
    ('W3C', 'World Wide Web Consortium'),
    ('WSN', 'Wireless Sensor Network')
]

results = []
for abbr, expansion in abbreviations:
    clean_exp = re.sub(r'\(.*?\)', '', expansion).strip()
    
    if abbr == 'F1':
        abbr_pattern = re.compile(r'\bF1\b|F1-score|F1 score', re.IGNORECASE)
    elif abbr == 'ER':
        abbr_pattern = re.compile(r'\bER\b|E-R\b|Entity-Relationship', re.IGNORECASE)
    else:
        abbr_pattern = re.compile(rf'\b{re.escape(abbr)}\b')
        
    exp_pattern = re.compile(re.escape(clean_exp), re.IGNORECASE)
    
    abbr_matches = []
    exp_matches = []
    
    for loc, pos, txt in texts:
        m1 = list(abbr_pattern.finditer(txt))
        m2 = list(exp_pattern.finditer(txt))
        if m1:
            abbr_matches.append((loc, pos, len(m1), txt.strip()))
        if m2:
            exp_matches.append((loc, pos, len(m2), txt.strip()))
            
    total_abbr = sum(m[2] for m in abbr_matches)
    total_exp = sum(m[2] for m in exp_matches)
    
    results.append({
        'abbr': abbr,
        'expansion': expansion,
        'abbr_count': total_abbr,
        'exp_count': total_exp,
        'total': total_abbr + total_exp,
        'abbr_matches': abbr_matches,
        'exp_matches': exp_matches
    })

print(f"{'Abbr':<6} | {'Full Expansion':<38} | {'Abbr Count':<10} | {'Exp Count':<10} | {'Status'}")
print('-' * 80)
for r in results:
    status = 'USED' if r['total'] > 0 else 'NOT USED (UNNECESSARY)'
    print(f"{r['abbr']:<6} | {r['expansion']:<38} | {r['abbr_count']:<10} | {r['exp_count']:<10} | {status}")

print("\n" + "="*80)
print("DETAILS FOR UNUSED OR LOW FREQUENCY ABBREVIATIONS:")
print("="*80)
for r in results:
    if r['total'] == 0:
        print(f"\n[UNNECESSARY / ZERO OCCURRENCES] {r['abbr']} - {r['expansion']}")
    elif r['total'] <= 2:
        print(f"\n[LOW USAGE: {r['total']}] {r['abbr']} - {r['expansion']}")
        for loc, pos, count, txt in (r['abbr_matches'] + r['exp_matches'])[:3]:
            print(f"   -> {loc} {pos}: {txt[:120]}...")
