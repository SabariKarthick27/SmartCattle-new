import docx
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r'c:\Project\SmartCattle new\QR MINI PROJECT REPORT.docx')

print("=== CHECKING ALL TABLES FOR RUMINATION / ACTIVITY ===")
for t_idx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            t = cell.text.lower()
            if 'rumination' in t or 'activity' in t or 'rumen' in t:
                # ignore workflow "activity diagram"
                if 'activity diagram' in t:
                    continue
                print(f"Table {t_idx+1} (R{r_idx}, C{c_idx}): {cell.text.strip()[:100]}")

print("\n=== CHECKING ALL PARAGRAPHS FOR RUMINATION / ACTIVITY ===")
for i, p in enumerate(doc.paragraphs):
    t = p.text.lower()
    if 'rumination' in t or 'activity' in t or 'rumen' in t:
        if 'activity diagram' in t:
            continue
        print(f"P[{i}]: {p.text.strip()}\n")
