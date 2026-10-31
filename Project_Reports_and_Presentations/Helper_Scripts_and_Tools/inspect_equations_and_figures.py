import docx

doc = docx.Document(r'c:\Project\SmartCattle new\QR MINI PROJECT REPORT.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")

eq_indices = [175, 177, 182, 184, 189, 191, 196, 199, 204, 206, 211, 213, 221, 226, 233, 235]
for i in eq_indices:
    if i < len(doc.paragraphs):
        p = doc.paragraphs[i]
        drawings = p._element.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing')
        omml = p._element.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath')
        print(f"P[{i}]: text='{p.text.strip()}', runs={len(p.runs)}, drawings={len(drawings)}, omml={len(omml)}")

print("\n--- ALL DRAWINGS / FIGURES IN DOCUMENT ---")
for i, p in enumerate(doc.paragraphs):
    drawings = p._element.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing')
    if drawings:
        print(f"P[{i}] has {len(drawings)} drawing(s). Text: '{p.text[:60]}'")
