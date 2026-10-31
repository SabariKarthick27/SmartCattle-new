import docx
import os
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r'c:\Project\SmartCattle new\QR MINI PROJECT REPORT.docx')

print("="*60)
print("FIGURE CAPTIONS AND ASSOCIATED PARAGRAPHS")
print("="*60)
for i, p in enumerate(doc.paragraphs):
    if "figure " in p.text.lower() or "fig." in p.text.lower() or "fig " in p.text.lower():
        print(f"P[{i}]: {p.text.strip()}")

print("\n" + "="*60)
print("TABLE CAPTIONS AND HEADINGS")
print("="*60)
for i, p in enumerate(doc.paragraphs):
    if "table " in p.text.lower():
        print(f"P[{i}]: {p.text.strip()}")
