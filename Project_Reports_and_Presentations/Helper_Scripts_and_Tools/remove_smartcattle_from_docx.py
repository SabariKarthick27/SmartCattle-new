import docx
import shutil
import sys
import re

file_path = r'c:\Project\SmartCattle new\SmartCattle_Project_Report_Aligned.docx'
backup_path = r'c:\Project\SmartCattle new\SmartCattle_Project_Report_Aligned.docx.bak'

# Read docx
doc = docx.Document(file_path)

replacements = [
    ("The proposed technique introduces 'SmartCattle' (QR-Cattle-Health)—an", "The proposed technique introduces a QR-based cattle health management system—an"),
    ("The proposed technique introduces 'SmartCattle' (QR-Cattle-Health)–an", "The proposed technique introduces a QR-based cattle health management system–an"),
    ("The proposed technique introduces 'SmartCattle' (QR-Cattle-Health)", "The proposed technique introduces a QR-based cattle health management system"),
    ("'SmartCattle' (QR-Cattle-Health)", "the proposed QR-based cattle management system"),
    ("SmartCattle (Mahalanobis CHDI)", "Proposed System (Mahalanobis CHDI)"),
    ("Comprehensive System Architecture of SmartCattle Multi-Tier Framework", "Comprehensive System Architecture of Proposed Multi-Tier Framework"),
    ("The SmartCattle architecture is structured", "The proposed system architecture is structured"),
    ("the SmartCattle system is organized", "the proposed system is organized"),
    ("The operational methodology of SmartCattle", "The operational methodology of the proposed system"),
    ("The mathematical foundation of SmartCattle", "The mathematical foundation of the proposed system"),
    ("SmartCattle is implemented as a modern", "The proposed system is implemented as a modern"),
    ("deployed SmartCattle platform.", "deployed platform."),
    ("interface screens of SmartCattle,", "interface screens of the proposed system,"),
    ("Day 18. SmartCattle triggered", "Day 18. The proposed system triggered"),
    ("Table 5.2 compares SmartCattle against", "Table 5.2 compares the proposed system against"),
    ("In contrast, SmartCattle normalizes", "In contrast, the proposed system normalizes"),
    ("detection performance of SmartCattle on", "detection performance of the proposed system on"),
    ("The SmartCattle project engineered", "The proposed project engineered"),
    ("2.4 days, SmartCattle offers a practical,", "2.4 days, the proposed system offers a practical,"),
    ("The SmartCattle project is directly", "The proposed project is directly"),
    ("impacts of SmartCattle are detailed", "impacts of the proposed system are detailed"),
    ("ketosis, SmartCattle prevents", "ketosis, the proposed system prevents"),
    ("primary objective, SmartCattle generates", "primary objective, the proposed system generates"),
    ("details of the SmartCattle platform.", "details of the proposed platform."),
    ("The SmartCattle platform is specifically", "The proposed platform is specifically"),
    # General fallback replacements
    ("SmartCattle platform", "proposed platform"),
    ("SmartCattle system", "proposed system"),
    ("SmartCattle project", "proposed project"),
    ("SmartCattle", "the proposed system")
]

def replace_in_run(run):
    modified = False
    for old, new in replacements:
        if old in run.text:
            run.text = run.text.replace(old, new)
            modified = True
    return modified

mod_count = 0
for i, p in enumerate(doc.paragraphs):
    for r in p.runs:
        if re.search(r'smart\s*cattle', r.text, re.I):
            if "https://github.com" in r.text:
                continue
            if replace_in_run(r):
                mod_count += 1
                print(f"Replaced in Paragraph {i}: {r.text[:80]}...")
            else:
                old_text = r.text
                r.text = re.sub(r'SmartCattle', 'the proposed system', r.text, flags=re.I)
                if r.text != old_text:
                    mod_count += 1
                    print(f"Regex replaced in Paragraph {i}: {r.text[:80]}...")

for ti, t in enumerate(doc.tables):
    for r_idx, row in enumerate(t.rows):
        for c_idx, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                for r in p.runs:
                    if re.search(r'smart\s*cattle', r.text, re.I):
                        if replace_in_run(r):
                            mod_count += 1
                            print(f"Replaced in Table {ti} R{r_idx}C{c_idx}: {r.text}...")
                        else:
                            old_text = r.text
                            r.text = re.sub(r'SmartCattle', 'Proposed System', r.text, flags=re.I)
                            if r.text != old_text:
                                mod_count += 1
                                print(f"Regex replaced in Table {ti} R{r_idx}C{c_idx}: {r.text}...")

# Try saving
try:
    doc.save(file_path)
    print(f"Successfully saved to {file_path}")
except PermissionError:
    alt_path = r'c:\Project\SmartCattle new\SmartCattle_Project_Report_Aligned_Updated.docx'
    doc.save(alt_path)
    print(f"File was locked by Microsoft Word. Saved updated copy to {alt_path}")
    sys.exit(2)
