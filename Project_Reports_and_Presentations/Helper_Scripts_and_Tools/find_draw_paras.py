import docx

doc = docx.Document(r'Project_Report_Final_20_10_26 (1).docx')
for i, p in enumerate(doc.paragraphs):
    has_draw = any('drawing' in r._element.xml for r in p.runs)
    if has_draw:
        prev_txt = doc.paragraphs[i-1].text.strip() if i > 0 else ""
        next_txt = doc.paragraphs[i+1].text.strip() if i+1 < len(doc.paragraphs) else ""
        print(f"P{i} [DRAWING]: prev='{prev_txt[:50]}' | next='{next_txt[:50]}'")
