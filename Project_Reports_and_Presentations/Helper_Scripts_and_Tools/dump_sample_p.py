import docx

doc_sample = docx.Document('Patent Sample (1).docx')
print(f"Total paragraphs: {len(doc_sample.paragraphs)}")
with open('sample_paragraphs_dump.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc_sample.paragraphs):
        f.write(f"P{i:03d} [{p.style.name}] (align={p.alignment}): {p.text}\n")
print("Saved sample_paragraphs_dump.txt")
