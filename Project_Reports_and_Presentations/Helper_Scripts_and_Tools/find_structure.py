import docx

doc = docx.Document("SmartCattle_Mini_Project_Report_v4.docx")

# Let's inspect paragraphs, tables, images and their exact layout blocks
blocks = []
current_page = 1
is_prelim = True

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

for i, p in enumerate(doc.paragraphs):
    has_break = any("<w:br" in r._r.xml and 'w:type="page"' in r._r.xml for r in p.runs)
    txt = p.text.strip()
    if has_break or "w:pageBreakBefore" in p._p.xml:
        print(f"--- PAGE BREAK before/at P{i}: {txt[:45]} ---")
    elif txt.startswith("CHAPTER") or txt.startswith("REFERENCES") or txt.startswith("APPENDIX") or txt in ["BONAFIDE CERTIFICATE", "DECLARATION", "ACKNOWLEDGEMENT", "ABSTRACT", "TABLE OF CONTENTS", "LIST OF TABLES", "LIST OF FIGURES", "LIST OF ABBREVIATIONS"]:
        print(f"P{i} [{p.style.name}]: {txt[:50]}")
