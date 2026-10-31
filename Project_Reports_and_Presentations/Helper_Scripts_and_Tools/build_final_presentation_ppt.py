import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE_TYPE
import os

def apply_text_styling(paragraph, lead_text, body_text, lead_size=12.5, body_size=12.5, lead_bold=True):
    paragraph.text = ""
    r1 = paragraph.add_run()
    r1.text = lead_text
    r1.font.name = "Arial"
    r1.font.size = Pt(lead_size)
    r1.font.bold = lead_bold
    r1.font.color.rgb = RGBColor(15, 23, 42) # DARK_BLUE
    
    if body_text:
        r2 = paragraph.add_run()
        r2.text = body_text
        r2.font.name = "Arial"
        r2.font.size = Pt(body_size)
        r2.font.color.rgb = RGBColor(51, 65, 85) # BODY_COLOR

def format_cell(cell, text, bold=False, size=10, color=RGBColor(51, 65, 85)):
    cell.text = text
    tf = cell.text_frame
    if tf.paragraphs and tf.paragraphs[0].runs:
        r = tf.paragraphs[0].runs[0]
        r.font.name = "Arial"
        r.font.bold = bold
        r.font.size = Pt(size)
        r.font.color.rgb = color

def build_presentation():
    template_path = r"c:\Project\SmartCattle new\MINI PROJECT PPT - REVIEW 2 (1).pptx"
    
    prs = Presentation(template_path)
    print(f"Loaded template presentation with {len(prs.slides)} slides.")

    DARK_BLUE = RGBColor(15, 23, 42)
    TEAL_ACCENT = RGBColor(13, 148, 136)
    BODY_COLOR = RGBColor(51, 65, 85)
    
    # -------------------------------------------------------------------------
    # SLIDE 1: Title & Team Details
    # -------------------------------------------------------------------------
    s1 = prs.slides[0]
    for shape in s1.shapes:
        if shape.name == 'Title 1' and shape.has_text_frame:
            tf = shape.text_frame
            tf.text = ""
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = "SONA COLLEGE OF TECHNOLOGY, SALEM\nDepartment of Information Technology"
            r.font.name = "Arial"
            r.font.size = Pt(14)
            r.font.bold = True
            r.font.color.rgb = DARK_BLUE

        elif shape.name == 'Subtitle 2' and shape.has_text_frame:
            tf = shape.text_frame
            tf.text = ""
            p1 = tf.paragraphs[0]
            p1.alignment = PP_ALIGN.CENTER
            r1 = p1.add_run()
            r1.text = "QR-BASED CATTLE MANAGEMENT SYSTEM"
            r1.font.name = "Arial"
            r1.font.size = Pt(24)
            r1.font.bold = True
            r1.font.color.rgb = DARK_BLUE
            
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            r2 = p2.add_run()
            r2.text = "Smart Health Monitoring & Automated Veterinary Alerts"
            r2.font.name = "Arial"
            r2.font.size = Pt(14)
            r2.font.italic = True
            r2.font.color.rgb = TEAL_ACCENT

        elif shape.name == 'TextBox 3' and shape.has_text_frame:
            tf = shape.text_frame
            tf.text = ""
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = "Project Guide:\n"
            r.font.name = "Arial"
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = DARK_BLUE
            
            r_g = p.add_run()
            r_g.text = "Mrs. J. Deepika,\nAssistant Professor,\nDepartment of IT"
            r_g.font.name = "Arial"
            r_g.font.size = Pt(12)
            r_g.font.color.rgb = BODY_COLOR

        elif shape.name == 'TextBox 4' and shape.has_text_frame:
            tf = shape.text_frame
            tf.text = ""
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = "Project Team Members:\n"
            r.font.name = "Arial"
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = DARK_BLUE
            
            members = [
                "Sabari Karthick A (61782323110093)",
                "Sivanarasimman S (61782323110103)",
                "Sujeeth Y R (61782323110107)"
            ]
            r_m = p.add_run()
            r_m.text = "\n".join(members)
            r_m.font.name = "Arial"
            r_m.font.size = Pt(12)
            r_m.font.color.rgb = BODY_COLOR

    # -------------------------------------------------------------------------
    # SLIDE 2: Abstract
    # -------------------------------------------------------------------------
    s2 = prs.slides[1]
    for shape in s2.shapes:
        if shape.name == 'Title 1' and shape.has_text_frame:
            shape.text_frame.text = "Abstract"
            shape.text_frame.paragraphs[0].font.name = "Arial"
            shape.text_frame.paragraphs[0].font.size = Pt(24)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = DARK_BLUE
        elif shape.name == 'Content Placeholder 2' and shape.has_text_frame:
            tf = shape.text_frame
            tf.text = ""
            bullets = [
                ("Low-Cost QR Identification: ", "Uses simple weatherproof QR ear tags instead of expensive wearable electronic sensors."),
                ("Digital Cattle Records: ", "Stores cattle profiles, milk yield, feed intake, vaccinations, and treatment history in one web app."),
                ("Personalized Health Tracking: ", "Learns the normal health baseline for each individual cow instead of using generic breed averages."),
                ("Early Disease Alerts: ", "Detects health drops (CHDI score) and alerts farmers and veterinarians 24 to 48 hours before severe illness."),
                ("Seamless Vet Coordination: ", "Automatically assigns sick cattle cases to nearby veterinarians for prompt treatment.")
            ]
            for i, (lead, body) in enumerate(bullets):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.space_after = Pt(10)
                apply_text_styling(p, "• " + lead, body, lead_size=13.5, body_size=13.5)

    # -------------------------------------------------------------------------
    # SLIDE 3: Objectives
    # -------------------------------------------------------------------------
    s3 = prs.slides[2]
    for shape in s3.shapes:
        if shape.name == 'Title 1' and shape.has_text_frame:
            shape.text_frame.text = "Objectives of the Project"
            shape.text_frame.paragraphs[0].font.name = "Arial"
            shape.text_frame.paragraphs[0].font.size = Pt(24)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = DARK_BLUE
        elif shape.name == 'Content Placeholder 2' and shape.has_text_frame:
            tf = shape.text_frame
            tf.text = ""
            obj_bullets = [
                ("Replace Paper Records: ", "Create an easy digital platform for farm and cattle management."),
                ("Fast QR Code Identification: ", "Enable instant cattle lookup by scanning ear tags using any smartphone camera."),
                ("Complete Medical History: ", "Track daily milk yield, feeding, vaccinations, and veterinary treatments in one secure place."),
                ("Personalized Health Baseline: ", "Monitor each cow's individual normal state to prevent false alarms."),
                ("Early Anomaly Detection: ", "Calculate a Composite Health Deviation Index (CHDI) to spot sickness early."),
                ("Automated Vet Alerts: ", "Automatically notify veterinarians when a cow shows signs of illness."),
                ("Affordable for Small Farms: ", "Deliver a simple, mobile-friendly system with zero on-animal hardware cost.")
            ]
            for i, (lead, body) in enumerate(obj_bullets):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.space_after = Pt(8)
                apply_text_styling(p, "• " + lead, body, lead_size=13, body_size=13)

    # -------------------------------------------------------------------------
    # SLIDE 4: Proposed System / Methodology
    # -------------------------------------------------------------------------
    s4 = prs.slides[3]
    for shape in s4.shapes:
        if shape.name == 'Title 1' and shape.has_text_frame:
            shape.text_frame.text = "Proposed System / Methodology"
            shape.text_frame.paragraphs[0].font.name = "Arial"
            shape.text_frame.paragraphs[0].font.size = Pt(24)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = DARK_BLUE
        elif shape.name == 'Content Placeholder 2' and shape.has_text_frame:
            tf = shape.text_frame
            tf.text = ""
            meth_bullets = [
                ("1. Registration & Tagging: ", "Farmer registers cattle and attaches a unique printed QR ear tag."),
                ("2. Daily Data Logging: ", "Farmer scans the QR tag with a smartphone and enters daily milk yield and feed intake."),
                ("3. Baseline Learning: ", "The system learns the cow's normal health patterns over a 14-day period."),
                ("4. Health Scoring (CHDI): ", "Compares new entries with the baseline to calculate a Health Score (0 to 100)."),
                ("5. Automated Triage: ", "Normal (≥80), Watch (60–79), and At-Risk (<60). Alerts are triggered automatically if health drops."),
                ("6. Vet Check & Treatment: ", "Veterinarian inspects the cow, adds treatment details, and confirms the diagnosis.")
            ]
            for i, (lead, body) in enumerate(meth_bullets):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.space_after = Pt(10)
                apply_text_styling(p, "• " + lead, body, lead_size=13.5, body_size=13.5)

    # -------------------------------------------------------------------------
    # SLIDE 5: Architecture / System Diagram
    # -------------------------------------------------------------------------
    s5 = prs.slides[4]
    for shape in list(s5.shapes):
        if shape.name == 'Title 1' and shape.has_text_frame:
            shape.text_frame.text = "System Architecture & Workflow"
            shape.text_frame.paragraphs[0].font.name = "Arial"
            shape.text_frame.paragraphs[0].font.size = Pt(24)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = DARK_BLUE
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE or 'Picture' in shape.name:
            sp_elem = shape._element
            sp_elem.getparent().remove(sp_elem)

    pro_fig_path = r"c:\Project\SmartCattle new\patent_figures\figure_master_pro_no_numbers.png"
    if os.path.exists(pro_fig_path):
        left = Inches(1.2)
        top = Inches(1.5)
        width = Inches(10.9)
        height = Inches(5.6)
        s5.shapes.add_picture(pro_fig_path, left, top, width, height)

    # -------------------------------------------------------------------------
    # SLIDE 6: Module Presentation
    # -------------------------------------------------------------------------
    s6 = prs.slides[5]
    for shape in s6.shapes:
        if shape.name == 'Title 1' and shape.has_text_frame:
            shape.text_frame.text = "Core Modules of the System"
            shape.text_frame.paragraphs[0].font.name = "Arial"
            shape.text_frame.paragraphs[0].font.size = Pt(24)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = DARK_BLUE
        elif shape.name == 'Content Placeholder 2' and shape.has_text_frame:
            tf = shape.text_frame
            tf.text = ""
            modules_list = [
                ("1. User Management: ", "Separate logins and permissions for Admin, Farmer, and Veterinarian."),
                ("2. Farm & Cattle Records: ", "Manage farm profiles, cow details, breed, age, and lactation stages."),
                ("3. QR Code Tagging: ", "Generate and scan unique QR codes for instant animal lookup."),
                ("4. Daily Health Logging: ", "Simple mobile form to log milk yield, feed intake, and activity."),
                ("5. Health Analysis Engine (CHDI): ", "Compares daily records against the cow's own baseline."),
                ("6. Smart Alert System: ", "Sends instant notifications when health drops or check-ups are due."),
                ("7. Veterinary Case Management: ", "Assigns sick cattle tickets directly to regional veterinarians."),
                ("8. Medical & Vaccination Records: ", "Tracks prescribed medicines, treatment logs, and vaccination dates.")
            ]
            for i, (lead, body) in enumerate(modules_list):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.space_after = Pt(6)
                apply_text_styling(p, "• " + lead, body, lead_size=12.5, body_size=12.5)

    # -------------------------------------------------------------------------
    # SLIDE 7: Module 1 — User, Farm & Cattle Management (ACTUAL PROJECT DETAILS)
    # -------------------------------------------------------------------------
    s7 = prs.slides[6]
    for shape in s7.shapes:
        if shape.name == 'Title 1' and shape.has_text_frame:
            shape.text_frame.text = "Module 1: User, Farm & Cattle Management"
            shape.text_frame.paragraphs[0].font.name = "Arial"
            shape.text_frame.paragraphs[0].font.size = Pt(24)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = DARK_BLUE
        
        elif shape.name == 'Table 2' and shape.has_table:
            t = shape.table
            format_cell(t.rows[0].cells[0], "User Roles & Access Control", bold=True, size=11, color=DARK_BLUE)
            format_cell(t.rows[1].cells[0], "• Admin: Platform logs & settings\n• Farmer: Farm herd & health logs", size=9.5)
            format_cell(t.rows[2].cells[0], "• Vet: Multi-farm health oversight", size=9.5)

        elif shape.name == 'Table 3' and shape.has_table:
            t = shape.table
            format_cell(t.rows[0].cells[0], "Farm Management", bold=True, size=11, color=DARK_BLUE)
            format_cell(t.rows[1].cells[0], "• Register farm location & capacity", size=9.5)
            format_cell(t.rows[2].cells[0], "• Assign responsible veterinarians", size=9.5)
            format_cell(t.rows[3].cells[0], "• View all cattle under each farm", size=9.5)

        elif shape.name == 'Table 4' and shape.has_table:
            t = shape.table
            format_cell(t.rows[0].cells[0], "Cattle Digital Profile", bold=True, size=11, color=DARK_BLUE)
            format_cell(t.rows[1].cells[0], "• Cattle ID (e.g. COW-001), breed, age", size=9.5)
            format_cell(t.rows[2].cells[0], "• Weight & lactation stage (lactating/dry)", size=9.5)
            format_cell(t.rows[3].cells[0], "• Live status: Healthy, Watch, At-Risk", size=9.5)

    print("Slide 7 updated.")

    # -------------------------------------------------------------------------
    # SLIDE 8: Module 2 — QR Code Generation & Optical Scanning (ACTUAL PROJECT DETAILS)
    # -------------------------------------------------------------------------
    s8 = prs.slides[7]
    for shape in s8.shapes:
        if shape.name == 'Title 1' and shape.has_text_frame:
            shape.text_frame.text = "Module 2: QR Code Generation & Tagging"
            shape.text_frame.paragraphs[0].font.name = "Arial"
            shape.text_frame.paragraphs[0].font.size = Pt(24)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = DARK_BLUE
        elif shape.name == 'TextBox 2' and shape.has_text_frame:
            tf = shape.text_frame
            tf.text = ""
            p1 = tf.paragraphs[0]
            apply_text_styling(p1, "• Weatherproof 2D QR Ear-Tags: ", "Every cow receives a unique printable QR tag linked to its digital cloud record. Eliminates expensive RFID tags with zero on-animal electronic sensors.", lead_size=12.5, body_size=12.5)
            
            p2 = tf.add_paragraph()
            p2.space_before = Pt(6)
            apply_text_styling(p2, "• Instant Phone Camera Scan (/scan): ", "Farmers scan the ear tag using standard smartphone cameras (no app install needed) to immediately view milk records, vaccination status, and log daily observations.", lead_size=12.5, body_size=12.5)

    print("Slide 8 updated.")

    # -------------------------------------------------------------------------
    # SLIDE 9: Module 3 — Health Monitoring & CHDI Analytics (ACTUAL PROJECT DETAILS)
    # -------------------------------------------------------------------------
    s9 = prs.slides[8]
    for shape in s9.shapes:
        if shape.name == 'Title 1' and shape.has_text_frame:
            shape.text_frame.text = "Module 3: Health Monitoring & CHDI Analytics"
            shape.text_frame.paragraphs[0].font.name = "Arial"
            shape.text_frame.paragraphs[0].font.size = Pt(24)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = DARK_BLUE
        elif shape.name == 'Content Placeholder 2' and shape.has_text_frame:
            tf = shape.text_frame
            tf.text = ""
            p0 = tf.paragraphs[0]
            p0.text = "Daily Parameters Logged:"
            p0.font.name = "Arial"
            p0.font.bold = True
            p0.font.size = Pt(13)
            p0.font.color.rgb = DARK_BLUE
            
            items = [
                "• Daily Milk Yield (Litres per milking)",
                "• Daily Feed Intake (kg) & Feed Type",
                "• Rumination Time & Activity Score",
                "• Core Body Temperature (°C) & Weight"
            ]
            for itm in items:
                p = tf.add_paragraph()
                p.text = itm
                p.font.name = "Arial"
                p.font.size = Pt(11.5)
                p.font.color.rgb = BODY_COLOR

        elif shape.name == 'Content Placeholder 3' and shape.has_text_frame:
            tf = shape.text_frame
            tf.text = ""
            p0 = tf.paragraphs[0]
            p0.text = "3-Stage CHDI Health Scoring:"
            p0.font.name = "Arial"
            p0.font.bold = True
            p0.font.size = Pt(13)
            p0.font.color.rgb = DARK_BLUE
            
            items = [
                "• 14-Day Baseline: Learns normal levels per cow",
                "• Deviation Score (0–100): Combines all vitals",
                "• Normal (≥80): Routine healthy logging",
                "• Watch (60–79): Advisory alert to farmer",
                "• At-Risk (<60): Urgent alert to veterinarian"
            ]
            for itm in items:
                p = tf.add_paragraph()
                p.text = itm
                p.font.name = "Arial"
                p.font.size = Pt(11.5)
                p.font.color.rgb = BODY_COLOR

    print("Slide 9 updated.")

    # -------------------------------------------------------------------------
    # SLIDE 10: Module 4 — Clinical Alerts & Veterinary Triage (ACTUAL PROJECT DETAILS)
    # -------------------------------------------------------------------------
    s10 = prs.slides[9]
    for shape in s10.shapes:
        if shape.name == 'Title 1' and shape.has_text_frame:
            shape.text_frame.text = "Module 4: Clinical Alerts & Veterinary Triage"
            shape.text_frame.paragraphs[0].font.name = "Arial"
            shape.text_frame.paragraphs[0].font.size = Pt(24)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = DARK_BLUE
        elif shape.name == 'Content Placeholder 2' and shape.has_text_frame:
            tf = shape.text_frame
            tf.text = ""
            p0 = tf.paragraphs[0]
            p0.text = "Automated Clinical Triage:"
            p0.font.name = "Arial"
            p0.font.bold = True
            p0.font.size = Pt(13)
            p0.font.color.rgb = DARK_BLUE
            
            items = [
                "• Instant ticket dispatch for At-Risk cows",
                "• Highlights top root causes (e.g. Milk Drop 54%)",
                "• Tracks emergency cases by urgency level",
                "• Regional disease outbreak radar for multi-farms"
            ]
            for itm in items:
                p = tf.add_paragraph()
                p.text = itm
                p.font.name = "Arial"
                p.font.size = Pt(11.5)
                p.font.color.rgb = BODY_COLOR

        elif shape.name == 'Content Placeholder 3' and shape.has_text_frame:
            tf = shape.text_frame
            tf.text = ""
            p0 = tf.paragraphs[0]
            p0.text = "Veterinary Records & Feedback:"
            p0.font.name = "Arial"
            p0.font.bold = True
            p0.font.size = Pt(13)
            p0.font.color.rgb = DARK_BLUE
            
            items = [
                "• Doctor records diagnosis (Mastitis, Ketosis)",
                "• Prescribes medicines, dosages & visit notes",
                "• CONFIRM / OVERRIDE feedback retrains model",
                "• Vaccination tracking (FMD, Anthrax, Brucellosis)"
            ]
            for itm in items:
                p = tf.add_paragraph()
                p.text = itm
                p.font.name = "Arial"
                p.font.size = Pt(11.5)
                p.font.color.rgb = BODY_COLOR

    print("Slide 10 updated.")

    # -------------------------------------------------------------------------
    # SLIDE 11: References
    # -------------------------------------------------------------------------
    s11 = prs.slides[10]
    for shape in s11.shapes:
        if shape.name == 'Title 1' and shape.has_text_frame:
            shape.text_frame.text = "References"
            shape.text_frame.paragraphs[0].font.name = "Arial"
            shape.text_frame.paragraphs[0].font.size = Pt(24)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = DARK_BLUE
        elif shape.name == 'Content Placeholder 2' and shape.has_text_frame:
            tf = shape.text_frame
            tf.text = ""
            refs = [
                ("[1] M. S. Mahmud et al., ", "“A systematic review on deep learning for precision cattle farming,” Computers & Electronics in Agriculture, 2021."),
                ("[2] J. Bumbálek, J. Pytel, and J. Císař, ", "“Dairy cow identification methods based on optical tags,” Computers & Electronics in Agriculture, 2023."),
                ("[3] X. Xu, B. Xiong, D. Liu, and H. Song, ", "“Non-contact technologies for cattle monitoring,” Computers & Electronics in Agriculture, 2026."),
                ("[4] C. J. Rutten, A. G. Velthuis, and H. Hogeveen, ", "“Sensors to support health management on dairy farms,” Journal of Dairy Science, 2013."),
                ("[5] D. Lovarelli, J. Bacenetti, and M. Guarino, ", "“Precision livestock farming in dairy cattle,” Journal of Cleaner Production, 2020.")
            ]
            for i, (tag, body) in enumerate(refs):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.space_after = Pt(8)
                apply_text_styling(p, tag, body, lead_size=12, body_size=12)
        elif shape.name == 'TextBox 3':
            sp = shape._element
            sp.getparent().remove(sp)

    # -------------------------------------------------------------------------
    # SLIDE 12: Conclusion & Future Scope
    # -------------------------------------------------------------------------
    s12 = prs.slides[11]
    for shape in s12.shapes:
        if shape.name == 'Title 1' and shape.has_text_frame:
            shape.text_frame.text = "Conclusion & Future Work"
            shape.text_frame.paragraphs[0].font.name = "Arial"
            shape.text_frame.paragraphs[0].font.size = Pt(24)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = DARK_BLUE
        elif shape.name == 'Content Placeholder 2' and shape.has_text_frame:
            tf = shape.text_frame
            tf.text = ""
            concl_bullets = [
                ("Affordable Cattle Tracking: ", "Provides smart health monitoring using low-cost QR ear tags without expensive wearable sensors."),
                ("Early Sickness Detection: ", "Identifies health drops (CHDI) 1 to 2 days earlier, allowing timely treatment before milk yield suffers."),
                ("Better Vet Coordination: ", "Connects farmers directly with veterinarians for rapid case response and clear treatment history."),
                ("Future Scope: ", "Adding offline mode for rural areas with low network connectivity, and voice assistance in Tamil and Hindi.")
            ]
            for i, (lead, body) in enumerate(concl_bullets):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.space_after = Pt(12)
                apply_text_styling(p, "• " + lead, body, lead_size=14, body_size=14)

    target_files = [
        r"c:\Project\SmartCattle new\SmartCattle_Mini_Project_Review_2_Simplified.pptx",
        r"c:\Project\SmartCattle new\SmartCattle_Mini_Project_Review_2_Final.pptx",
        r"c:\Project\SmartCattle new\MINI PROJECT PPT - REVIEW 2 (1).pptx"
    ]
    
    for tf in target_files:
        try:
            prs.save(tf)
            print(f"Successfully saved to: {tf}")
        except Exception as e:
            print(f"Could not save to {tf} (locked): {e}")

if __name__ == "__main__":
    build_presentation()
