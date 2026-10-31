import sys
import docx
import os
import re

doc_path = r'c:\Project\SmartCattle new\QR MINI PROJECT REPORT.docx'
doc = docx.Document(doc_path)

out_file = r'c:\Project\SmartCattle new\audit_report_findings.txt'

with open(out_file, 'w', encoding='utf-8') as f:
    f.write(f"DOCUMENT AUDIT FOR: {doc_path}\n")
    f.write(f"Total Paragraphs: {len(doc.paragraphs)}\n")
    f.write(f"Total Tables: {len(doc.tables)}\n\n")

    f.write("="*80 + "\n")
    f.write("1. STRUCTURAL HEADINGS AUDIT\n")
    f.write("="*80 + "\n")
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if re.match(r'^(CHAPTER|\d+(\.\d+)*|[A-Z\s]{4,})', txt) and len(txt) < 100:
            f.write(f"P[{i}]: {txt}\n")

    f.write("\n" + "="*80 + "\n")
    f.write("2. SUSPICIOUS / GARBLED / TEMPLATE LEFTOVER STRINGS\n")
    f.write("="*80 + "\n")
    suspicious_keywords = [
        'student', 'saltlens', 'venom', 'todo', 'fixme', 'lorem', 'template',
        'sample', 'placeholder', 'error', 'undefined', 'nan', '\ufffd', '??', 'xx'
    ]
    for i, p in enumerate(doc.paragraphs):
        txt = p.text
        for kw in suspicious_keywords:
            if kw in txt.lower():
                f.write(f"P[{i}] (keyword: '{kw}'): {txt}\n\n")

    f.write("\n" + "="*80 + "\n")
    f.write("3. ALL TABLES AUDIT\n")
    f.write("="*80 + "\n")
    for t_idx, table in enumerate(doc.tables):
        f.write(f"\n--- TABLE {t_idx+1} ({len(table.rows)} rows x {len(table.columns)} cols) ---\n")
        for r_idx, row in enumerate(table.rows):
            row_txt = [c.text.strip().replace('\n', ' ') for c in row.cells]
            # remove duplicates caused by merged cells
            dedup = []
            for cell in row_txt:
                if not dedup or cell != dedup[-1]:
                    dedup.append(cell)
            f.write(f"  R{r_idx}: {' | '.join(dedup)}\n")

    f.write("\n" + "="*80 + "\n")
    f.write("4. CODE IMPLEMENTATION SNIPPETS IN REPORT (SECTION 4.13)\n")
    f.write("="*80 + "\n")
    in_code = False
    for i, p in enumerate(doc.paragraphs):
        if "4.13" in p.text:
            in_code = True
        if "4.14" in p.text:
            in_code = False
        if in_code:
            f.write(f"P[{i}]: {p.text}\n")

    f.write("\n" + "="*80 + "\n")
    f.write("5. CHAPTER 5 / RESULTS AND DISCUSSION AUDIT\n")
    f.write("="*80 + "\n")
    in_ch5 = False
    for i, p in enumerate(doc.paragraphs):
        if "5.1" in p.text or "CHAPTER 5" in p.text or "RESULTS AND DISCUSSION" in p.text:
            in_ch5 = True
        if "6.1" in p.text or "CHAPTER 6" in p.text:
            in_ch5 = False
        if in_ch5:
            f.write(f"P[{i}]: {p.text}\n")

print("Audit script completed. Output written to audit_report_findings.txt")
