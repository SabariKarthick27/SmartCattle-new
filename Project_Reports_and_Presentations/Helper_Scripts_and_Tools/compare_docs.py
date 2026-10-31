import docx

doc_sample = docx.Document('Patent Sample (1).docx')
doc_cattle = docx.Document('QR_Based_Cattle_Management_Patent_Disclosure.docx')

print("=== SAMPLE SECTIONS ===")
for i, s in enumerate(doc_sample.sections):
    print(f"Sec {i}: top={s.top_margin.pt}, bot={s.bottom_margin.pt}, left={s.left_margin.pt}, right={s.right_margin.pt}")

print("\n=== COMPARING PARAGRAPHS COUNT ===")
print("Sample:", len(doc_sample.paragraphs))
print("Cattle:", len(doc_cattle.paragraphs))

print("\n=== COMPARING TABLES COUNT ===")
print("Sample:", len(doc_sample.tables))
print("Cattle:", len(doc_cattle.tables))

for t_idx, (st, ct) in enumerate(zip(doc_sample.tables, doc_cattle.tables)):
    print(f"Table {t_idx}: Sample={len(st.rows)}x{len(st.columns)}, Cattle={len(ct.rows)}x{len(ct.columns)}")
