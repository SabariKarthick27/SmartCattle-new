import docx
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r'c:\Project\SmartCattle new\QR MINI PROJECT REPORT.docx')

print("="*70)
print("EXACT PARAGRAPH BY PARAGRAPH MAPPING OF ALL 24 IMAGES")
print("="*70)

for i, p in enumerate(doc.paragraphs):
    blips = p._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
    if blips:
        print(f"\n[Paragraph {i}] has {len(blips)} image(s):")
        if p.text.strip():
            print(f"  Self text: '{p.text.strip()}'")
        for j in range(max(0, i-2), min(len(doc.paragraphs), i+4)):
            t = doc.paragraphs[j].text.strip()
            if any(k in t.lower() for k in ['figure', 'fig.', 'equation', 'table']):
                print(f"  --> Context P[{j}]: '{t}'")
