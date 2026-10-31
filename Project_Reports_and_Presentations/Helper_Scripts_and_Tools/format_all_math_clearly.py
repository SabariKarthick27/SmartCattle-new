import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os
import shutil

COLOR_BLACK = RGBColor(0, 0, 0)

def set_run_font(run, font_name="Times New Roman", size_pt=12, bold=False, italic=False, color_rgb=(0,0,0)):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)

def add_p_to_doc(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6, line_spacing=1.15, bold=False, italic=False, size_pt=12):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    r = p.add_run(text)
    set_run_font(r, font_name="Times New Roman", size_pt=size_pt, bold=bold, italic=italic)
    return p

def format_all_math_in_document(docx_path):
    doc = docx.Document(docx_path)
    print(f"Loaded {docx_path} with {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables.")

    # 1. Clean Table 10 and Table 11
    if len(doc.tables) > 10:
        t10 = doc.tables[10] # Table 4.3 Severity Bands
        for row in t10.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.text = (p.text
                        .replace("$D_M$", "DM")
                        .replace("$D_M < 0.70$", "DM < 0.70")
                        .replace("$0.70 \\le D_M < 1.45$", "0.70 ≤ DM < 1.45")
                        .replace("$1.45 \\le D_M < 2.65$", "1.45 ≤ DM < 2.65")
                        .replace("$D_M \\ge 2.65$", "DM ≥ 2.65")
                        .replace("$", ""))
                    for r in p.runs:
                        set_run_font(r, font_name="Times New Roman", size_pt=12)

    if len(doc.tables) > 11:
        t11 = doc.tables[11] # Table 4.4 Multi-Signal Features
        for row in t11.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.text = (p.text
                        .replace("($x_k$)", "(xk)")
                        .replace("($x_1$)", "(x₁)")
                        .replace("($x_2$)", "(x₂)")
                        .replace("($x_3$)", "(x₃)")
                        .replace("($x_4$)", "(x₄)")
                        .replace("($x_5$)", "(x₅)")
                        .replace("$", ""))
                    for r in p.runs:
                        set_run_font(r, font_name="Times New Roman", size_pt=12)

    # 2. Fix inline text paragraphs across the entire document
    for p in doc.paragraphs:
        t = p.text
        if "$x_1$" in t or "($x_1$)" in t:
            p.text = (t
                .replace("($x_1$)", "(x₁)")
                .replace("($x_2$)", "(x₂)")
                .replace("($x_3$)", "(x₃)")
                .replace("($x_4$)", "(x₄)")
                .replace("($x_5$)", "(x₅)")
                .replace("$x_1$", "x₁")
                .replace("$x_2$", "x₂")
                .replace("$x_3$", "x₃")
                .replace("$x_4$", "x₄")
                .replace("$x_5$", "x₅"))
            for r in p.runs: set_run_font(r, font_name="Times New Roman", size_pt=12)

        if "$D_M$" in t and "\\frac" not in t:
            p.text = p.text.replace("$D_M$", "DM (Mahalanobis distance)")
            for r in p.runs: set_run_font(r, font_name="Times New Roman", size_pt=12)

        if "$S < -2.0$" in t:
            p.text = p.text.replace("$k=7$", "k = 7").replace("($S < -2.0$)", "(slope S < −2.0)")
            for r in p.runs: set_run_font(r, font_name="Times New Roman", size_pt=12)

        if "$N \\ge 7$" in t or "$N \\ge 7$" in t or "$N_min" in t or "$N_{min}" in t:
            p.text = p.text.replace("$N \\ge 7$", "N ≥ 7 days").replace("$N_{min}=7$", "N_min = 7 days").replace("$N_{min} = 7$", "N_min = 7 days")
            for r in p.runs: set_run_font(r, font_name="Times New Roman", size_pt=12)

        if "$|W| = 14$" in t:
            p.text = p.text.replace("$|W| = 14$", "|W| = 14 days").replace("$\\lambda_{reg} = 10^{-3}$", "λ_reg = 10⁻³").replace("$\\Sigma$", "Σ")
            for r in p.runs: set_run_font(r, font_name="Times New Roman", size_pt=12)

        if "\\frac{TP}{TP + FN}" in t:
            # Replace evaluation metrics block with beautifully formatted, clean text
            p.text = (
                "System evaluation follows standard statistical metrics:\n"
                "• Sensitivity (Recall): Sensitivity = TP / (TP + FN), measuring the proportion of true clinical anomalies detected.\n"
                "• Specificity: Specificity = TN / (TN + FP), measuring resistance against false alarms on healthy baseline days.\n"
                "• Precision: Precision = TP / (TP + FP), measuring the proportion of generated alerts corresponding to genuine clinical distress.\n"
                "• F1-Score: F1-Score = 2 × (Precision × Recall) / (Precision + Recall), the harmonic mean of precision and recall.\n"
                "• Diagnostic Lead Time: Advance warning lead time (2.4 days on average) between initial sub-clinical CHDI deviation and observable clinical symptoms."
            )
            for r in p.runs: set_run_font(r, font_name="Times New Roman", size_pt=12)

    # 3. Restructure Section 4.7 Mathematical Formulation (P173 & P175)
    # Find P173 and P175
    idx_p173 = -1
    idx_p175 = -1
    for i, p in enumerate(doc.paragraphs):
        if "The mathematical foundation of the proposed system executes across three analytical stages" in p.text:
            idx_p173 = i
        if "Stage 3: Temporal Trend, Explainability, and Uncertainty Estimation:" in p.text:
            idx_p175 = i

    print(f"Found Section 4.7 Stage 1 & 2 at P{idx_p173}, Stage 3 at P{idx_p175}")

    if idx_p173 != -1:
        # We replace the raw LaTeX text of P173 with clean intro and insert equations
        p_intro = doc.paragraphs[idx_p173]
        p_intro.text = (
            "The mathematical foundation of the proposed system executes across three analytical stages as formalized in Section IV of the IEEE paper:\n\n"
            "Stage 1: Individual Rolling Baseline & Standardized Deviations:\n"
            "Rather than comparing an animal against arbitrary population-wide static thresholds, Stage 1 establishes an individualized historical baseline over a 14-day observation window (|W| = 14). For each physiological parameter k (where k ∈ {1, 2, 3, 4, 5}), the baseline mean and sample variance are computed using Equation (4.1) and Equation (4.2)."
        )
        for r in p_intro.runs: set_run_font(r, font_name="Times New Roman", size_pt=12)

        # Build replacement paragraphs before Table 4.3
        # In python-docx, to insert elements cleanly before a target paragraph, we can insert XML siblings
        # or use insert_paragraph_before
        target_p = doc.paragraphs[idx_p173 + 1] # The paragraph right after P173 (which is Table 4.3 title)

        def add_eq_before(target_para, img_path, eq_num, eq_title, where_bullets):
            # 1. Equation image paragraph
            p_img = target_para.insert_paragraph_before()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(2)
            p_img.paragraph_format.keep_with_next = True
            r_img = p_img.add_run()
            r_img.add_picture(img_path, width=Inches(4.8))

            # 2. Equation caption paragraph
            p_cap = target_para.insert_paragraph_before()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(1)
            p_cap.paragraph_format.space_after = Pt(3)
            p_cap.paragraph_format.keep_with_next = True
            r_cap = p_cap.add_run(f"Equation ({eq_num}): {eq_title}")
            set_run_font(r_cap, font_name="Times New Roman", size_pt=11.5, bold=True)

            # 3. Where explanation
            p_where = target_para.insert_paragraph_before()
            p_where.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_where.paragraph_format.space_before = Pt(2)
            p_where.paragraph_format.space_after = Pt(8)
            p_where.paragraph_format.line_spacing = 1.15
            r_w_title = p_where.add_run("Where:\n")
            set_run_font(r_w_title, font_name="Times New Roman", size_pt=11, bold=True, italic=True)
            for bullet in where_bullets:
                r_b = p_where.add_run(f"• {bullet}\n")
                set_run_font(r_b, font_name="Times New Roman", size_pt=11)

        def add_text_before(target_para, text):
            p = target_para.insert_paragraph_before()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(text)
            set_run_font(r, font_name="Times New Roman", size_pt=12)
            return p

        # Equation 4.1
        add_eq_before(target_p, "patent_figures/equations/eq_1.png", "4.1", "Rolling Baseline Mean Estimation", [
            "μ_{i,k}(t) represents the individualized 14-day rolling historical baseline mean for cow i across parameter k.",
            "|W| denotes the 14-day observation window length (|W| = 14 consecutive daily observations).",
            "x_{i,k}(t - m) is the historical daily measurement logged on day (t - m)."
        ])

        # Equation 4.2
        add_eq_before(target_p, "patent_figures/equations/eq_2.png", "4.2", "Rolling Baseline Sample Variance", [
            "σ_{i,k}²(t) is the unbiased sample variance capturing animal i's natural biological variation.",
            "(|W| - 1) is Bessel's correction factor for sample variance degrees of freedom (df = 13).",
            "(x_{i,k}(t - m) - μ_{i,k}(t))² is the squared deviation from the rolling baseline mean."
        ])

        # Explanatory text for Z-score
        add_text_before(target_p, 
            "When a new daily measurement is recorded, the raw value is normalized into a dimensionless standardized Z-score deviation using Equation (4.3). A regularization constant ε = 10⁻⁴ is included in the denominator to ensure numerical stability and prevent division-by-zero."
        )

        # Equation 4.3
        add_eq_before(target_p, "patent_figures/equations/eq_3.png", "4.3", "Standardized Individual Deviation (Z-Score)", [
            "z_{i,k}(t) is the standardized Z-score deviation of parameter k from cow i's personal baseline.",
            "x_{i,k}(t) is the newly observed measurement (e.g., daily milk yield, feed intake, rumination).",
            "ε = 10⁻⁴ is the numerical stabilization constant preventing zero-variance division."
        ])

        # Stage 2 Intro
        add_text_before(target_p,
            "Stage 2: Multivariate Covariance & Mahalanobis CHDI Formulation:\n"
            "Because bovine disorders affect multiple physiological systems simultaneously (e.g., subclinical mastitis causes concurrent drops in milk and feed intake alongside elevated core body temperature), Stage 2 models cross-variable dependencies using the regularized empirical covariance matrix Σ in Equation (4.4)."
        )

        # Equation 4.4
        add_eq_before(target_p, "patent_figures/equations/eq_4.png", "4.4", "Regularized Empirical Covariance Matrix", [
            "Σ(t) is the 5 × 5 empirical covariance matrix capturing physiological correlations across all 5 telemetry features.",
            "z_i(t - m) is the 5-dimensional standardized deviation vector at observation step (t - m).",
            "λ_reg · I is Tikhonov ridge shrinkage regularization (λ_reg = 10⁻³), guaranteeing positive-definiteness and stable inversion."
        ])

        # Explanatory text for Mahalanobis
        add_text_before(target_p,
            "The multi-signal anomaly distance is calculated in covariance-weighted space using the multivariate Mahalanobis distance in Equation (4.5). This distance quantifies the multi-dimensional departure of current vitals from the animal's personal physiological centroid."
        )

        # Equation 4.5
        add_eq_before(target_p, "patent_figures/equations/eq_5.png", "4.5", "Multivariate Mahalanobis Anomaly Distance", [
            "D_M(z_i) is the multivariate Mahalanobis anomaly distance.",
            "z_i is the normalized 5-dimensional deviation vector: [z_weight, z_feed, z_milk, z_temp, z_activity]ᵀ.",
            "Σ⁻¹ is the inverted regularized covariance matrix weighting correlated signals."
        ])

        # Explanatory text for CHDI mapping
        add_text_before(target_p,
            "To provide an intuitive metric for farmers and clinicians, the raw Mahalanobis distance D_M is mapped to a bounded Composite Health Deviation Index (CHDI) on a 0 to 100 scale using the exponential decay formulation in Equation (4.6)."
        )

        # Equation 4.6
        add_eq_before(target_p, "patent_figures/equations/eq_6.png", "4.6", "Bounded Composite Health Deviation Index (CHDI)", [
            "CHDI(z_i) is the continuous health score bounded between 100 (Optimal Health) and 0 (Severe Critical Distress).",
            "λ = 0.35 is the calibrated exponential decay factor.",
            "Table 4.3 defines the resulting clinical severity bands, distance boundaries, and veterinary escalation protocols."
        ])

    # Now handle Stage 3: Replace P175
    # Find P175 again (since paragraphs shifted)
    idx_p175_new = -1
    for i, p in enumerate(doc.paragraphs):
        if "Stage 3: Temporal Trend, Explainability, and Uncertainty Estimation:" in p.text:
            idx_p175_new = i
            break

    if idx_p175_new != -1:
        p_st3 = doc.paragraphs[idx_p175_new]
        p_st3.text = (
            "Stage 3: Temporal Trend, Explainability, and Uncertainty Estimation:\n"
            "Stage 3 enhances clinical utility through three algorithmic decision mechanisms: evaluating 7-day temporal trends, attributing feature contributions, and enforcing uncertainty gating."
        )
        for r in p_st3.runs: set_run_font(r, font_name="Times New Roman", size_pt=12)

        target_st3 = doc.paragraphs[idx_p175_new + 1] # Paragraph right after P175 (before Figure 4.3)

        # Equation 4.7
        add_eq_before(target_st3, "patent_figures/equations/eq_7.png", "4.7", "7-Day Temporal Trend Slope (OLS Regression)", [
            "S represents the 7-day health trajectory slope evaluated across the last k = 7 consecutive CHDI scores.",
            "t is the time step index (t ∈ {1, 2, ..., 7}) with midpoint t̄ = 4.",
            "A negative slope (S < −2.0) indicates rapid health deterioration, escalating alert priority before static bounds are breached."
        ])

        # Equation 4.8
        add_eq_before(target_st3, "patent_figures/equations/eq_8.png", "4.8", "Permutation Feature Importance Attribution", [
            "C_k is the percentage contribution (%) of physiological parameter k to the total anomaly alert.",
            "D_M(z) is the full Mahalanobis distance across all features.",
            "D_M(z_{-k}) is the counterfactual Mahalanobis distance evaluated with variable k held at baseline.",
            "Delivers transparent explainability so clinicians know which specific vitals triggered the alert."
        ])

        # Equation 4.9
        add_eq_before(target_st3, "patent_figures/equations/eq_9.png", "4.9", "Data-Sufficiency Uncertainty Gate", [
            "Conf(N) is the confidence coefficient scaling between 0.0 and 1.0.",
            "N is the count of consecutive valid daily observation logs recorded for the cow.",
            "N_min = 7 is the minimum logging threshold required before issuing definitive clinical alerts, preventing false alarms on newly registered cows."
        ])

    # Save to temp file and overwrite
    temp_out = "temp_formatted_math.docx"
    doc.save(temp_out)
    shutil.copy2(temp_out, docx_path)
    os.remove(temp_out)
    print(f"Successfully formatted all mathematical variables and equations in {docx_path}!")

if __name__ == "__main__":
    for f in ["SmartCattle_Project_Report_Aligned_MASTER.docx", "SmartCattle_Project_Report_Aligned.docx", "SmartCattle_Project_Report.docx"]:
        if os.path.exists(f):
            try:
                format_all_math_in_document(f)
            except Exception as e:
                print(f"Error on {f}: {e}")
