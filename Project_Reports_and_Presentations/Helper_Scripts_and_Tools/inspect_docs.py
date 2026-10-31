import docx

sample = docx.Document('Patent Sample (1).docx')
cattle = docx.Document('QR_Based_Cattle_Management_Patent_Disclosure.docx')

print("=== CHECKING DRAWINGS IN SAMPLE ===")
for p_idx, p in enumerate(sample.paragraphs):
    for r_idx, r in enumerate(p.runs):
        if 'drawing' in r._element.xml:
            print(f"Drawing in P{p_idx} R{r_idx}: text='{p.text}'")

for t_idx, tbl in enumerate(sample.tables):
    for r_idx, row in enumerate(tbl.rows):
        for c_idx, cell in enumerate(row.cells):
            for p_idx, p in enumerate(cell.paragraphs):
                for run_idx, r in enumerate(p.runs):
                    if 'drawing' in r._element.xml:
                        print(f"Drawing in TBL{t_idx} R{r_idx} C{c_idx} P{p_idx}: text='{cell.text[:40]}'")

print("\n=== CHECKING DRAWINGS IN CATTLE ===")
for p_idx, p in enumerate(cattle.paragraphs):
    for r_idx, r in enumerate(p.runs):
        if 'drawing' in r._element.xml:
            print(f"Drawing in P{p_idx} R{r_idx}: text='{p.text}'")

for t_idx, tbl in enumerate(cattle.tables):
    for r_idx, row in enumerate(tbl.rows):
        for c_idx, cell in enumerate(row.cells):
            for p_idx, p in enumerate(cell.paragraphs):
                for run_idx, r in enumerate(p.runs):
                    if 'drawing' in r._element.xml:
                        print(f"Drawing in TBL{t_idx} R{r_idx} C{c_idx} P{p_idx}: text='{cell.text[:40]}'")
