import docx
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r'c:\Project\SmartCattle new\QR MINI PROJECT REPORT.docx')
for idx, table in enumerate(doc.tables):
    for r_i, row in enumerate(table.rows):
        cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
        for c_i, cell in enumerate(cells):
            if cell == '':
                # check if signature row in certificate/declaration
                if idx in [0, 1, 2, 3]:
                    continue
                print(f"Table {idx+1} Row {r_i} Col {c_i} is EMPTY! Row text: {' | '.join(cells)}")
            if '\ufffd' in cell:
                print(f"Table {idx+1} Row {r_i} Col {c_i} has REPLACEMENT CHAR (\ufffd): {cell}")

print("\n--- Checking all paragraphs for \ufffd ---")
for i, p in enumerate(doc.paragraphs):
    if '\ufffd' in p.text:
        print(f"P[{i}] has \ufffd: {p.text}")
