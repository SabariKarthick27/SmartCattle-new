import docx

doc_sample = docx.Document('Patent Sample (1).docx')
p = doc_sample.paragraphs[72]
print(f"P72 runs in template: {len(p.runs)}")
for i, r in enumerate(p.runs):
    print(f"  Run {i}: '{r.text}'")

p29 = doc_sample.paragraphs[29]
print(f"P29 runs in template: {len(p29.runs)}")
for i, r in enumerate(p29.runs):
    print(f"  Run {i}: '{r.text}'")
