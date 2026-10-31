import docx

doc = docx.Document(r'Project_Report_Final_20_10_26 (1).docx')
for i, p in enumerate(doc.paragraphs):
    xml = p._p.xml
    if 'pic:pic' in xml or 'w:drawing' in xml:
        # find preceding or following text
        surrounding = [doc.paragraphs[j].text.strip() for j in range(max(0, i-2), min(len(doc.paragraphs), i+3)) if doc.paragraphs[j].text.strip()]
        print(f"Paragraph {i}:")
        for s in surrounding:
            print(f"   {s[:80]}".encode('ascii', errors='replace').decode('ascii'))
