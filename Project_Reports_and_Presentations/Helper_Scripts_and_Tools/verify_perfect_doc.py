import docx

doc = docx.Document('QR_Based_Cattle_Management_Patent_Disclosure.docx')
print(f"Total paragraphs: {len(doc.paragraphs)}")

print("\n=== VERIFYING PARAGRAPH SAMPLES ===")
check_indices = [2, 3, 4, 5, 27, 29, 30, 31, 55, 56, 57, 58, 59, 60, 71, 72, 73, 74, 75, 76, 77, 78, 79, 80, 83, 84, 85, 86, 93, 107, 123, 134, 147, 159, 171, 182, 183, 184, 185, 186, 187, 188, 189, 190]
for i in check_indices:
    p = doc.paragraphs[i]
    f_name = p.runs[0].font.name if p.runs else None
    f_size = p.runs[0].font.size.pt if (p.runs and p.runs[0].font.size) else None
    print(f"P{i:03d} [{p.style.name}] (font={f_name}, size={f_size}):\n   {p.text}")

print("\n=== VERIFYING TABLES SAMPLES ===")
for t_idx, tbl in enumerate(doc.tables):
    print(f"\n--- TABLE {t_idx} ---")
    for r_idx, row in enumerate(tbl.rows):
        for c_idx, cell in enumerate(row.cells):
            t = cell.text.replace('\n', ' ')
            if len(t) > 60:
                t = t[:60] + "..."
            print(f"  R{r_idx}C{c_idx}: {t}")
