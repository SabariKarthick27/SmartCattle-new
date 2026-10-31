import docx
import re

doc = docx.Document(r'c:\Project\SmartCattle new\SmartCattle_Project_Report_PERFECT_MATH.docx')

texts = []
for p in doc.paragraphs:
    if 'LIST OF ABBREVIATIONS' in p.text:
        continue
    if p.text.strip():
        texts.append(('paragraph', p.text))
for t_idx, t in enumerate(doc.tables):
    if t_idx == 7:
        continue
    for r in t.rows:
        for c in r.cells:
            if c.text.strip():
                texts.append(('table', c.text))

user_list = [
    ('API', 'Application Programming Interface'),
    ('BLE', 'Bluetooth Low Energy'),
    ('BRD', 'Bovine Respiratory Disease'),
    ('CHDI', 'Composite Health Deviation Index'),
    ('EWMA', 'Exponentially Weighted Moving Average'),
    ('F1', 'Harmonic Mean of Precision and Recall'),
    ('IoT', 'Internet of Things'),
    ('LF', 'Low Frequency (RFID)'),
    ('ML', 'Machine Learning'),
    ('MQTT', 'Message Queuing Telemetry Transport'),
    ('OCR', 'Optical Character Recognition'),
    ('OLS', 'Ordinary Least Squares'),
    ('ORM', 'Object-Relational Mapping'),
    ('PLF', 'Precision Livestock Farming'),
    ('PWA', 'Progressive Web Application'),
    ('QR', 'Quick Response Code'),
    ('RBAC', 'Role-Based Access Control'),
    ('REST', 'Representational State Transfer'),
    ('RFID', 'Radio-Frequency Identification'),
    ('THI', 'Temperature-Humidity Index'),
    ('UHF', 'Ultra-High Frequency (RFID)'),
    ('UI', 'User Interface'),
    ('W3C', 'World Wide Web Consortium'),
    ('WSN', 'Wireless Sensor Network')
]

print(f"{'Abbr':<6} | {'Abbr In Text':<12} | {'Exp In Text':<12} | {'Verdict'}")
print("-" * 65)

remove_list = []
expansion_only = []
keep_list = []

for abbr, exp in user_list:
    clean_exp = re.sub(r'\(.*?\)', '', exp).strip()
    
    if abbr == 'F1':
        p_abbr = re.compile(r'\bF1\b|F1-score|F1 score', re.IGNORECASE)
    else:
        p_abbr = re.compile(rf'\b{re.escape(abbr)}\b')
    p_exp = re.compile(re.escape(clean_exp), re.IGNORECASE)
    
    c_abbr = sum(len(p_abbr.findall(t[1])) for t in texts)
    c_exp = sum(len(p_exp.findall(t[1])) for t in texts)
    
    if c_abbr == 0 and c_exp == 0:
        verdict = "REMOVE (Zero mentions)"
        remove_list.append((abbr, exp))
    elif c_abbr == 0 and c_exp > 0:
        verdict = "OPTIONAL REMOVE (Only full phrase used)"
        expansion_only.append((abbr, exp, c_exp))
    else:
        verdict = f"KEEP (Abbr: {c_abbr}, Exp: {c_exp})"
        keep_list.append((abbr, exp, c_abbr))
        
    print(f"{abbr:<6} | {c_abbr:<12} | {c_exp:<12} | {verdict}")

print("\n--- SUMMARY ---")
print("MUST REMOVE (0 occurrences):")
for a, e in remove_list:
    print(f"  - {a} ({e})")

print("\nCONSIDER REMOVING (Expansion only, acronym never written):")
for a, e, c in expansion_only:
    print(f"  - {a} ({e}) -> phrase appears {c} times, but abbreviation '{a}' never appears")

print(f"\nKEEP ({len(keep_list)} items):")
for a, e, c in keep_list:
    print(f"  - {a} ({e})")
