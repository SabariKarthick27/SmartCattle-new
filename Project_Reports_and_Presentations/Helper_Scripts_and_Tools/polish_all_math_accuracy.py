import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import shutil
import os

def set_run_font(run, font_name="Times New Roman", size_pt=11, bold=False, italic=False, subscript=False, superscript=False, color_rgb=(0,0,0)):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.subscript = subscript
    run.font.superscript = superscript
    run.font.color.rgb = RGBColor(*color_rgb)
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)

def polish_all_math_accuracy():
    doc_path = "SmartCattle_Project_Report_PERFECT_MATH.docx"
    doc = docx.Document(doc_path)
    print(f"Loaded {doc_path} for deep mathematical polishing...")

    # 1. Polish Table 10: D_M subscripts
    if len(doc.tables) > 10:
        t10 = doc.tables[10]
        # Header
        c_hdr = t10.rows[0].cells[1]
        c_hdr.text = ""
        p = c_hdr.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p.add_run("Mahalanobis Distance (")
        set_run_font(r1, size_pt=12, bold=True)
        r2 = p.add_run("D")
        set_run_font(r2, size_pt=12, bold=True, italic=True)
        r3 = p.add_run("M")
        set_run_font(r3, size_pt=12, bold=True, italic=True, subscript=True)
        r4 = p.add_run(")")
        set_run_font(r4, size_pt=12, bold=True)

        # Rows 1 to 4
        thresholds = [
            ("DM < 0.70", [("D", True, False), ("M", True, True), (" < 0.70", False, False)]),
            ("0.70 ≤ DM < 1.45", [("0.70 ≤ ", False, False), ("D", True, False), ("M", True, True), (" < 1.45", False, False)]),
            ("1.45 ≤ DM < 2.65", [("1.45 ≤ ", False, False), ("D", True, False), ("M", True, True), (" < 2.65", False, False)]),
            ("DM ≥ 2.65", [("D", True, False), ("M", True, True), (" ≥ 2.65", False, False)])
        ]
        for r_idx, (_, parts) in enumerate(thresholds, start=1):
            cell = t10.rows[r_idx].cells[1]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for txt, is_it, is_sub in parts:
                r = p.add_run(txt)
                set_run_font(r, size_pt=12, italic=is_it, subscript=is_sub)

    # 2. Polish Table 11: Parameter (xₖ)
    if len(doc.tables) > 11:
        t11 = doc.tables[11]
        c_p = t11.rows[0].cells[0]
        c_p.text = ""
        p = c_p.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = p.add_run("Parameter (")
        set_run_font(r1, size_pt=12, bold=True)
        r2 = p.add_run("x")
        set_run_font(r2, size_pt=12, bold=True, italic=True)
        r3 = p.add_run("k")
        set_run_font(r3, size_pt=12, bold=True, italic=True, subscript=True)
        r4 = p.add_run(")")
        set_run_font(r4, size_pt=12, bold=True)

    # 3. Polish body paragraphs for perfect subscripts and typography
    for p in doc.paragraphs:
        t = p.text
        
        # Polish P199: λ_reg
        if "Tikhonov ridge shrinkage regularization (λ_reg" in t:
            p.text = ""
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run("• ")
            set_run_font(r, size_pt=11, bold=True)
            r = p.add_run("λ")
            set_run_font(r, size_pt=11, italic=True)
            r = p.add_run("reg")
            set_run_font(r, size_pt=11, italic=True, subscript=True)
            r = p.add_run(" · I = Tikhonov ridge shrinkage regularization (λ")
            set_run_font(r, size_pt=11)
            r = p.add_run("reg")
            set_run_font(r, size_pt=11, subscript=True)
            r = p.add_run(" = 10⁻³), guaranteeing positive-definite matrix inversion.")
            set_run_font(r, size_pt=11)

        # Polish P205: z_i and feature vector
        if "Normalized 5-dimensional deviation vector: [z_weight" in t or "zi = Normalized" in t:
            p.text = ""
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run("• ")
            set_run_font(r, size_pt=11, bold=True)
            r = p.add_run("z")
            set_run_font(r, size_pt=11, italic=True)
            r = p.add_run("i")
            set_run_font(r, size_pt=11, italic=True, subscript=True)
            r = p.add_run(" = Standardized 5-signal deviation vector across all monitored features [")
            set_run_font(r, size_pt=11)
            
            features = [("weight", "z"), ("feed", "z"), ("milk", "z"), ("temp", "z"), ("activity", "z")]
            for idx, (f_name, z_sym) in enumerate(features):
                r = p.add_run(z_sym)
                set_run_font(r, size_pt=11, italic=True)
                r = p.add_run(f_name)
                set_run_font(r, size_pt=11, subscript=True)
                if idx < len(features) - 1:
                    r = p.add_run(", ")
                    set_run_font(r, size_pt=11)
            r = p.add_run("]ᵀ.")
            set_run_font(r, size_pt=11)

        # Polish D_M in narrative text
        if "raw Mahalanobis distance D_M is mapped" in t:
            p.text = ""
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run("To provide an intuitive metric for farmers and clinicians, the raw Mahalanobis distance ")
            set_run_font(r, size_pt=12)
            r = p.add_run("D")
            set_run_font(r, size_pt=12, italic=True)
            r = p.add_run("M")
            set_run_font(r, size_pt=12, italic=True, subscript=True)
            r = p.add_run(" is mapped to a bounded Composite Health Deviation Index (CHDI) on a 0 to 100 scale using the exponential decay formulation in Equation (4.6).")
            set_run_font(r, size_pt=12)

        # Polish Section 4.9 preprocessing
        if "Tikhonov ridge regularization (λ_reg = 10⁻³)" in t:
            p.text = (
                "Raw field logs undergo standard validation and preprocessing:\n"
                "1. Range Verification: Biologically implausible values (e.g., negative milk yield or body temperature > 45°C) are rejected at form entry.\n"
                "2. 14-Day Rolling Window Isolation: For each cow, observations within the last 14 days (|W| = 14 days) are extracted to update baseline statistics.\n"
                "3. Regularized Inversion: Tikhonov ridge regularization (λreg = 10⁻³) is applied to covariance matrix Σ to prevent singularity under collinear telemetry."
            )
            for r in p.runs: set_run_font(r, font_name="Times New Roman", size_pt=12)

        # Polish Section 5.1 evaluation metrics block
        if "System evaluation follows standard statistical metrics:" in t and "Sensitivity = TP" in t:
            p.text = ""
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run("System evaluation follows standard statistical metrics:\n")
            set_run_font(r, size_pt=12, bold=True)
            
            metrics = [
                ("• Sensitivity (Recall)", "Sensitivity = TP / (TP + FN)", "Measures the proportion of true clinical disease anomalies successfully detected by the system."),
                ("• Specificity", "Specificity = TN / (TN + FP)", "Measures the system's resistance against false alarms during normal healthy baseline days."),
                ("• Precision", "Precision = TP / (TP + FP)", "Measures the proportion of generated alerts that correspond to genuine clinical distress."),
                ("• F1-Score", "F1-Score = 2 × (Precision × Recall) / (Precision + Recall)", "The harmonic mean balancing precision and sensitivity."),
                ("• Diagnostic Lead Time", "Lead Time = 2.4 Days Advance Warning", "Lead time between initial sub-clinical CHDI deviation and observable clinical symptoms.")
            ]
            for name, formula, desc in metrics:
                r_n = p.add_run(f"{name}: ")
                set_run_font(r_n, size_pt=11.5, bold=True)
                r_f = p.add_run(f"{formula}\n")
                set_run_font(r_f, size_pt=11.5, italic=True)
                r_d = p.add_run(f"   {desc}\n")
                set_run_font(r_d, size_pt=11)

    # Save polished file
    doc.save(doc_path)
    print(f"Deep mathematical polishing completed on {doc_path}!")

    # Synchronize to SmartCattle_Project_Report.docx
    try:
        shutil.copy2(doc_path, "SmartCattle_Project_Report.docx")
        print("Synchronized to SmartCattle_Project_Report.docx")
    except PermissionError:
        print("SmartCattle_Project_Report.docx is locked.")

    # Artifact copy
    art_path = r"C:\Users\SABARIKARTHICK\.gemini\antigravity-ide\brain\5ea3760d-61ec-4b9d-a301-3985c4c1874a\SmartCattle_Project_Report_PERFECT_MATH.docx"
    shutil.copy2(doc_path, art_path)
    print("Saved to artifacts directory.")

if __name__ == "__main__":
    polish_all_math_accuracy()
