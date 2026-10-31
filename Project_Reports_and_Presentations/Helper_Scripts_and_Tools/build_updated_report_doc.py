import docx
import shutil

src_path = r'Project_Report_Final_20_10_26 (1).docx'
dst_path = r'Project_Report_Final_Updated.docx'

shutil.copyfile(src_path, dst_path)
doc = docx.Document(dst_path)

# Let's inspect paragraphs around equations 4.4 to 4.9
# 1. Update text within existing paragraphs
for i, p in enumerate(doc.paragraphs):
    # Eq 4.4 text update
    if "5 × 5 empirical covariance matrix" in p.text:
        p.text = p.text.replace("5 × 5 empirical covariance matrix capturing dependencies across all 5 telemetry features.",
                                "4 × 4 empirical covariance matrix capturing dependencies across the 4 core telemetry features (body weight, feed intake, milk production, core temperature).")
        print("Updated Eq 4.4 text")
    
    # Eq 4.5 text update
    if "Standardized 5-signal deviation vector" in p.text:
        p.text = p.text.replace("Standardized 5-signal deviation vector across all monitored features [zweight, zfeed, zmilk, ztemp]ᵀ.",
                                "Standardized 4-signal deviation vector across all monitored features [zweight, zfeed, zmilk, ztemp]ᵀ.")
        print("Updated Eq 4.5 text")
        
    # Eq 4.6 bands update
    if "Clinical Bands: Healthy (80–100)" in p.text:
        p.text = p.text.replace("Clinical Bands: Healthy (80–100), Watch (60–79), Warning (40–59), and Critical (0–39).",
                                "Clinical Bands: Healthy (≥ 75), Watch (50–74), At-Risk (< 50), and Abstain (N < 7 or Ct < 0.40).")
        print("Updated Eq 4.6 bands")

    # Eq 4.8 text update
    if "Full Mahalanobis anomaly distance across all 5 telemetry features." in p.text:
        p.text = p.text.replace("all 5 telemetry features.",
                                "all 4 telemetry features (body weight, feed intake, milk production, core temperature).")
        print("Updated Eq 4.8 text")

# Now let's locate the paragraphs containing the equation images for 4.7, 4.8, 4.9
# For Eq 4.7: P223 is the drawing paragraph right after P222 ("Stage 3 enhances clinical utility...")
# For Eq 4.8: P228 is the drawing paragraph right after P227 ("• Clinical Rule: Negative slope...")
# For Eq 4.9: P233 is the drawing paragraph right after P232 ("• DM(z−k) = Counterfactual distance...")

# Let's find them dynamically by inspecting preceding text
p_eq7 = None
p_eq8 = None
p_eq9 = None

for i, p in enumerate(doc.paragraphs):
    has_pic = 'pic:pic' in p._p.xml or 'w:drawing' in p._p.xml
    if has_pic and i > 0:
        prev = doc.paragraphs[i-1].text.strip()
        if "Stage 3 enhances clinical utility" in prev:
            p_eq7 = p
            print(f"Found Eq 4.7 pic at P{i}")
        elif "Clinical Rule: Negative slope" in prev:
            p_eq8 = p
            print(f"Found Eq 4.8 pic at P{i}")
        elif "Counterfactual distance evaluated" in prev:
            p_eq9 = p
            print(f"Found Eq 4.9 pic at P{i}")

def set_para_font(p, font_name="Times New Roman", size_pt=12):
    for r in p.runs:
        r.font.name = font_name
        r.font.size = docx.shared.Pt(size_pt)

# 2. Insert narrative use case paragraphs before Eq 4.7, Eq 4.8, Eq 4.9
if p_eq7:
    intro7 = p_eq7.insert_paragraph_before(
        "To capture downward trajectory momentum and identify early insidious health degradation before bounds are breached, the rate of change across consecutive daily health scores is evaluated using the 7-day linear regression slope in Equation (4.7)."
    )
    intro7.paragraph_format.line_spacing = 1.5
    intro7.paragraph_format.space_after = docx.shared.Pt(6)
    set_para_font(intro7)
    print("Inserted use-case before Eq 4.7")

if p_eq8:
    intro8 = p_eq8.insert_paragraph_before(
        "To provide transparent diagnostic rationale for clinicians, multi-signal alert departures are decomposed into parameter-specific percentage contributions using the Leave-One-Covariate-Out (LOCO) permutation attribution in Equation (4.8)."
    )
    intro8.paragraph_format.line_spacing = 1.5
    intro8.paragraph_format.space_after = docx.shared.Pt(6)
    set_para_font(intro8)
    print("Inserted use-case before Eq 4.8")

if p_eq9:
    intro9 = p_eq9.insert_paragraph_before(
        "To prevent spurious false alarms during the initial onboarding of newly registered cattle with limited observation history, the data-sufficiency abstention gate is enforced using Equation (4.9)."
    )
    intro9.paragraph_format.line_spacing = 1.5
    intro9.paragraph_format.space_after = docx.shared.Pt(6)
    set_para_font(intro9)
    print("Inserted use-case before Eq 4.9")

doc.save(dst_path)
print(f"Successfully generated {dst_path}!")
